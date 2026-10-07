"""Readable, numeric single-location workbook from recorded model snapshots.

Export never fetches weather or reruns the visibility model. Missing diagnostics
remain blank; sunset and each method's independently chosen optimum stay apart.
"""

import math
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from hilal_visibility.models.crumey import B_FLOOR, Z_V, nL_to_cdm2, sr_to_arcmin2, cdm2_to_mag_arcsec2, mag_to_lux
from hilal_visibility.models.kastner import S10_TO_NL


NAVY, TEAL, PURPLE, PALE, INK = '17324D', '087E8B', '7657A5', 'EDF3F8', '22364A'
GREEN, RED, GREY = 'DFF1E5', 'FCE5E5', 'E7EDF2'
DECIMAL, SCIENTIFIC, TIME = '0.000000', '0.000000E+00', 'yyyy-mm-dd hh:mm:ss'
NA = 'Tidak tersedia'
NE_LABEL, TEL_LABEL = 'Mata telanjang', 'Teleskop'


def _number(value):
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)  # +inf/-inf are meaningful thresholds/margins, never 0.
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)  # Wall time; timezone is explicitly labelled.
    return value


def _cell(ws, row, col, value, fmt=DECIMAL, fill=None, bold=False):
    cell = ws.cell(row, col, _number(value))
    if isinstance(value, str):
        cell.data_type = 's'  # User text beginning with '=' must remain literal text.
    cell.font = Font(name='Calibri', size=11, color=INK, bold=bold)
    cell.alignment = Alignment(vertical='center', wrap_text=True,
                               horizontal='left' if isinstance(value, str) else 'right')
    cell.border = Border(bottom=Side(style='hair', color='D8E2EB'))
    cell.number_format = TIME if isinstance(value, datetime) else 'yyyy-mm-dd' if isinstance(value, date) else fmt
    if fill:
        cell.fill = PatternFill('solid', fgColor=fill)
    return cell


def _bar(ws, row, title, last=8, color=NAVY):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last)
    for col in range(1, last + 1):
        cell = ws.cell(row, col)
        cell.fill = PatternFill('solid', fgColor=color)
    cell = ws.cell(row, 1, title)
    cell.font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
    cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[row].height = 26


def _page(ws, title, subtitle, widths):
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_properties.tabColor = NAVY
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = 'Halaman &P / &N'
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    _bar(ws, 1, title, len(widths))
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(widths))
    _cell(ws, 2, 1, subtitle)
    ws.row_dimensions[2].height = 34


def _matrix(ws, title, subtitle):
    _page(ws, title, subtitle, [34, 19, 17, 23, 23, 23, 23, 64])
    for left, right, title, color in ((4, 5, 'SAAT SUNSET', TEAL), (6, 7, 'WAKTU OPTIMAL MASING-MASING', PURPLE)):
        ws.merge_cells(start_row=4, start_column=left, end_row=4, end_column=right)
        for col in range(left, right + 1):
            ws.cell(4, col).fill = PatternFill('solid', fgColor=color)
        cell = ws.cell(4, left, title)
        cell.font = Font(name='Calibri', bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    headers = ['Parameter', 'Simbol', 'Satuan', NE_LABEL, TEL_LABEL, NE_LABEL, TEL_LABEL, 'Rumus / keterangan']
    for col, label in enumerate(headers, 1):
        _cell(ws, 5, col, label, fill=PALE, bold=True)
    ws.row_dimensions[4].height = 30
    ws.freeze_panes = 'D6'
    ws.print_title_rows = '1:5'


def _scene_values(calc, scene, telescope=False):
    if not scene or scene.get('valid') is False:
        return {}
    if telescope and not calc.hasil.get('use_telescope', True):
        return {}
    trace = scene.get('model_trace') or {}
    phot = trace.get('photometry') or {}
    ne = trace.get('naked_eye') or {}
    tel = trace.get('telescope') or {}
    optics, threshold = tel.get('optics') or {}, tel.get('threshold') or {}
    coefficients = (tel.get('coefficients') if telescope else trace.get('naked_eye_coefficients')) or {}
    config = calc.hasil.get('tel_params') or {}
    suffix = 'tel' if telescope else 'ne'
    area = ne.get('A_sr')
    area_deg2 = phot.get('area_deg2')
    L, B = scene.get('luminansi_hilal_nl'), scene.get('sky_brightness_nl')
    L_app = scene.get('luminansi_hilal_tel_nl') if telescope else L
    B_app = scene.get('sky_brightness_tel_nl') if telescope else B
    C_th = scene.get('c_th_tel' if telescope else 'crumey_ne_C_th')
    C_obj, margin = scene.get('rasio_kontras_' + suffix), scene.get('delta_m_' + suffix)
    g = optics.get('surface_brightness_factor') if telescope else 1.0
    app_area = threshold.get('A_app_sr') if telescope else area
    if telescope and app_area is None and area is not None and config.get('magnification'):
        app_area = area * config['magnification']**2
    convert = lambda x: nL_to_cdm2(x) if x is not None else None
    L_cd, B_cd, L_app_cd, B_app_cd = map(convert, (L, B, L_app, B_app))
    delta_sky = threshold.get('delta_B_th') if telescope else (C_th * B_cd if C_th is not None and B_cd is not None else None)
    delta_app = threshold.get('delta_B_app_th') if telescope else delta_sky
    illumination = (1 + math.cos(math.radians(scene['phase_angle']))) / 2 if scene.get('phase_angle') is not None else None
    status = ('DI BAWAH HORIZON' if scene.get('moon_alt', 1) <= 0 else
              'TERLIHAT (MODEL)' if margin is not None and margin > 0 else 'TIDAK TERLIHAT (MODEL)')
    intrinsic = phot.get('intrinsic_luminance_nL')
    m_v, a_v = phot.get('M_v'), scene.get('extinction_mag_v')
    flux = L_cd * area if area is not None and L_cd is not None else None
    v = {
        **scene,
        'time': scene.get('waktu_local') or scene.get('sunset_local'),
        'age_hours': None, 'width_arcsec': scene['moon_width'] * 3600 if scene.get('moon_width') is not None else None,
        'illuminated_fraction': illumination, 'area_fraction': None,
        'area_deg2': area_deg2, 'area_arcmin2': ne.get('A_arcmin2'), 'area_sr': area,
        'M_v': m_v, 'L_star_s10': phot.get('L_star_s10'), 'intrinsic_nL': intrinsic,
        'intrinsic_cd': convert(intrinsic), 'phase_flux_lux': mag_to_lux(m_v) if m_v is not None else None,
        'M_v_atmosphere': m_v + a_v if m_v is not None and a_v is not None else None,
        'L_nL': L, 'L_cd': L_cd, 'B_nL': B, 'B_cd': B_cd, 'flux_lux': flux,
        'mu_sky': cdm2_to_mag_arcsec2(B_cd) if B_cd is not None and B_cd > 0 else None,
        'D_ap': config.get('aperture') if telescope else None,
        'D_s': config.get('central_obstruction') if telescope else None,
        'M': config.get('magnification') if telescope else 1.0,
        'p': optics.get('De') if telescope else None,
        'd': optics.get('exit_pupil') if telescope else None,
        'd_secondary': optics.get('secondary_exit_pupil') if telescope else None,
        'tau': optics.get('transmission') if telescope else 1.0,
        'admitted_fraction': optics.get('admitted_pupil_fraction') if telescope else 1.0,
        'Ft': optics.get('Ft') if telescope else 1.0,
        'g': g, 'area_app_sr': app_area,
        'area_app_arcmin2': sr_to_arcmin2(app_area) if app_area is not None else None,
        'L_app_nL': L_app, 'B_app_nL': B_app, 'L_app_cd': L_app_cd, 'B_app_cd': B_app_cd,
        'F': config.get('field_factor') if telescope else calc.hasil.get('F_naked'),
        'FT': math.sqrt(2) if telescope else 1.0, 'FM': 1.0,
        'phi': threshold.get('phi') if telescope else calc.hasil.get('F_naked'),
        'B_eval': coefficients.get('B_eval'), 'B_floor': B_FLOOR,
        'floor': coefficients.get('floor_applied'), 'curve': coefficients.get('curve'),
        'regime': coefficients.get('regime'),
        'achromatic_extrapolation': coefficients.get('achromatic_extrapolation'),
        'field_factor_comparison': scene.get('field_factor_comparison') if telescope else None,
        'R': coefficients.get('R'), 'C_inf': coefficients.get('C_inf'), 'q': coefficients.get('q'),
        'ricco_area_sr': coefficients.get('ricco_area_sr'),
        'C_obj': C_obj, 'C_th': C_th, 'delta_app': delta_app, 'delta_sky': delta_sky,
        'ratio': C_obj / C_th if C_obj is not None and C_th is not None and C_th > 0 else None,
        'margin': margin, 'log_margin': margin / 2.5 if margin is not None else None,
        'status': status,
        'gain': scene.get('threshold_difference_mag', scene.get('telescope_gain')) if telescope else None,
    }
    if area_deg2 is not None and scene.get('moon_semidiameter'):
        v['area_fraction'] = area_deg2 / (math.pi * scene['moon_semidiameter']**2)
    conjunction = calc.hasil.get('ijtima_utc')
    if conjunction and v['time']:
        v['age_hours'] = (v['time'] - conjunction).total_seconds() / 3600
    return v


# (key, label, symbol, unit, explanation, number format)
CHAIN = [
    '01 | GEOMETRI DAN WAKTU',
    ('time', 'Waktu evaluasi lokal', 't', 'zona waktu lokasi', 'Optimal NE dan teleskop dipilih secara independen.', TIME),
    ('age_hours', 'Umur Bulan sejak ijtima', 't − t_ijtima', 'jam', 'Selisih waktu evaluasi dan ijtima UTC.', DECIMAL),
    ('sun_alt', 'Altitude Matahari', 'h_sun', 'deg', 'Posisi dengan refraksi dinamis T/P.', DECIMAL),
    ('sun_az', 'Azimuth Matahari', 'Az_sun', 'deg', 'Ephemeris Skyfield / DE440s.', DECIMAL),
    ('moon_alt', 'Altitude Bulan', 'h_moon', 'deg', 'Posisi dengan refraksi dinamis T/P.', DECIMAL),
    ('moon_az', 'Azimuth Bulan', 'Az_moon', 'deg', 'Ephemeris Skyfield / DE440s.', DECIMAL),
    ('elongation', 'Elongasi toposentrik', 'E', 'deg', 'Separasi Matahari–Bulan; dipakai untuk luas model.', DECIMAL),
    ('phase_angle', 'Sudut fase', 'α', 'deg', 'Sudut fase astrometrik; dipakai untuk magnitudo.', DECIMAL),
    ('moon_distance_km', 'Jarak Bulan toposentrik', 'Δ', 'km', 'Jarak astrometrik pengamat–Bulan.', '0.000'),
    ('moon_semidiameter', 'Semidiameter Bulan', 'r', 'deg', 'asin(R_bulan / Δ).', DECIMAL),
    ('width_arcsec', 'Lebar sabit dari fase', 'W', 'arcsec', 'r × (1 + cos α); dikonversi dari derajat.', DECIMAL),
    ('illuminated_fraction', 'Fraksi iluminasi dari fase', 'k_phase', 'fraksi', '(1 + cos α) / 2; berbeda dari pendekatan luas berbasis E.', SCIENTIFIC),
    ('area_fraction', 'Fraksi luas model', 'k_area', 'fraksi', 'sin²(E/2); pendekatan geometri luas yang digunakan model.', SCIENTIFIC),
    ('area_deg2', 'Luas sabit model', 'D_luas', 'deg²', 'π r² sin²(E/2); D di rumus Kastner adalah luas.', SCIENTIFIC),
    ('area_arcmin2', 'Luas sabit model', 'A', 'arcmin²', 'D_luas × 3600.', SCIENTIFIC),
    ('area_sr', 'Sudut ruang sabit', 'A', 'sr', 'D_luas × (π/180)².', SCIENTIFIC),
    '02 | FOTOMETRI INTRINSIK — KASTNER',
    ('M_v', 'Magnitudo V intrinsik', 'M_v', 'mag', '0.026 α + 4×10⁻⁹ α⁴ − 12.73; phase law yang dipertahankan.', DECIMAL),
    ('phase_flux_lux', 'Fluks phase law sebelum atmosfer', 'I_V,phase', 'lux', 'Z_V × 10^(−0.4 M_v); phase law tidak validasi sendiri limit sabit tipis.', SCIENTIFIC),
    ('L_star_s10', 'Luminansi intrinsik Kastner', 'L*', 'S10', '10^[0.4(10 − M_v)] / D_luas; luas nol memberi sumber nol di implementasi.', SCIENTIFIC),
    ('intrinsic_nL', 'Luminansi intrinsik', 'L_intrinsik', 'nL', 'L* × faktor S10 → nL dari zero point V bersama.', SCIENTIFIC),
    ('intrinsic_cd', 'Luminansi intrinsik SI', 'L_intrinsik', 'cd/m²', 'Konversi nL → cd/m² yang sama dengan Crumey.', SCIENTIFIC),
    '03 | ATMOSFER DAN CAHAYA YANG TIBA DI PENGAMAT',
    ('rh', 'Kelembapan terkoreksi', 'RH', '%', 'Atmosfer pada waktu evaluasi; input Schaefer.', '0.00'),
    ('temperature', 'Suhu terkoreksi', 'T', '°C', 'Atmosfer pada waktu evaluasi; input Schaefer dan refraksi.', '0.00'),
    ('pressure', 'Tekanan udara', 'P', 'hPa', 'Dipakai untuk refraksi; bukan input eksplisit ekstingsi Schaefer.', '0.00'),
    ('k_v', 'Koefisien ekstingsi V', 'k_v', 'mag/airmass', 'k_R + k_A + k_O + k_W; bukan kehilangan total sepanjang LOS.', DECIMAL),
    ('extinction_mag_v', 'Ekstingsi total sepanjang LOS', 'A_V = DM[2]', 'mag', 'k_R X_G + k_A X_A + k_O X_O + k_W X_G; bukan margin Δm.', DECIMAL),
    ('transmission_v', 'Transmisi atmosfer V', 'T_V', 'fraksi', '10^(−0.4 A_V); diterapkan sekali pada luminansi intrinsik.', SCIENTIFIC),
    ('M_v_atmosphere', 'Magnitudo phase law setelah atmosfer', 'M_v + A_V', 'mag', 'Identitas fotometri untuk fluks langsung; bukan limiting magnitude.', DECIMAL),
    ('L_nL', 'Luminansi excess hilal', 'L', 'nL', 'L_intrinsik × T_V; tambahan cahaya hilal di atas latar.', SCIENTIFIC),
    ('L_cd', 'Luminansi excess hilal SI', 'ΔB_obj', 'cd/m²', 'L dalam SI; tidak mengurangi B_sky dari L.', SCIENTIFIC),
    ('flux_lux', 'Fluks hilal dari luas model', 'I_obs = L × A', 'lux', 'L [cd/m²] × A [sr]; uji identitas dengan fluks phase law × T_V untuk A > 0.', SCIENTIFIC),
    ('B_nL', 'Kecerlangan latar Schaefer', 'B_sky', 'nL', 'Atmosfer sudah masuk: night + min(twilight, daylight); tanpa moonlight.', SCIENTIFIC),
    ('B_cd', 'Kecerlangan latar SI', 'B_sky', 'cd/m²', 'Tidak dikalikan lagi dengan T_V.', SCIENTIFIC),
    ('mu_sky', 'Surface brightness langit V', 'μ_sky', 'mag/arcsec²', 'Dari B_sky dan zero point V yang sama; bukan magnitudo terintegrasi.', DECIMAL),
    '04 | OPTIK — TARGET DAN LATAR MENERIMA FAKTOR YANG SAMA',
    ('D_ap', 'Diameter aperture', 'D_ap', 'mm', 'D pada kontrak teleskop adalah diameter, bukan D_luas Kastner.', '0.00'),
    ('D_s', 'Diameter obstruksi pusat', 'D_s', 'mm', 'Diameter bayangan sekunder diperhitungkan pada pupil terpusat.', '0.00'),
    ('M', 'Pembesaran', 'M', '×', 'M = 1 untuk mata telanjang.', '0.00'),
    ('p', 'Diameter pupil mata yang digunakan', 'p', 'mm', 'Override jika tersedia; selain itu estimasi pupil gelap dari usia.', DECIMAL),
    ('d', 'Diameter exit pupil', 'd = D_ap/M', 'mm', 'Bukan diameter pupil mata.', DECIMAL),
    ('d_secondary', 'Diameter bayangan sekunder', 'd_s = D_s/M', 'mm', 'Pemotongan annulus oleh pupil dihitung eksplisit.', DECIMAL),
    ('tau', 'Transmisi total permukaan optik', 'τ = t₁ⁿ', 'fraksi', 'Teleskop: transmisi per permukaan dipangkatkan jumlah permukaan.', SCIENTIFIC),
    ('admitted_fraction', 'Fraksi pupil mata yang menerima cahaya', 'f_pupil', 'fraksi', 'max(min(d,p)² − d_s², 0) / p².', SCIENTIFIC),
    ('Ft', 'Faktor transmisi diagnostik', 'F_t', '—', '1 / [τ(1 − (D_s/D_ap)²)]; throughput aktif memakai g annular.', DECIMAL),
    ('g', 'Faktor luminansi optik aktif', 'g', 'fraksi', 'τ × f_pupil; L_app = gL dan B_app = gB. Untuk NE: g = 1.', SCIENTIFIC),
    ('area_app_sr', 'Sudut ruang tampak', 'A_app = M² A', 'sr', 'Luas aktual tampak tetap digunakan termasuk di bawah floor.', SCIENTIFIC),
    ('area_app_arcmin2', 'Luas tampak', 'A_app', 'arcmin²', 'Konversi A_app; asumsi seluruh sabit muat di medan pandang.', SCIENTIFIC),
    ('L_app_nL', 'Luminansi hilal setelah optik', 'L_app', 'nL', 'g × L.', SCIENTIFIC),
    ('B_app_nL', 'Latar setelah optik', 'B_app', 'nL', 'g × B_sky.', SCIENTIFIC),
    ('L_app_cd', 'Luminansi hilal tampak SI', 'L_app', 'cd/m²', 'Input increment yang diterima mata.', SCIENTIFIC),
    ('B_app_cd', 'Latar tampak SI', 'B_app', 'cd/m²', 'Latar aktual untuk kontras; tidak diganti floor.', SCIENTIFIC),
    '05 | PARAMETER DAN AMBANG CRUMEY',
    ('F', 'Residual visual field factor', 'F', '—', 'Laboratory scaling, pengamat, efek target/viewing residual; atmosfer, throughput, M dan FT/FM terpisah. Belum dikalibrasi khusus hilal.', DECIMAL),
    ('FT', 'Faktor monokular', 'F_T', '—', 'Teleskop √2; NE 1. Mengoreksi threshold, bukan luminansi.', DECIMAL),
    ('FM', 'Faktor tambahan optik/pengamat', 'F_M', '—', 'Saat ini 1; seeing, blur dan pembatasan medan belum dimodelkan.', DECIMAL),
    ('phi', 'Field factor efektif', 'φ = F F_T F_M', '—', 'Multiplier threshold aktif.', DECIMAL),
    ('curve', 'Kurva threshold yang digunakan', 'mode', '—', 'auto memakai combined di semua latar; tidak switching dengan label regime.', DECIMAL),
    ('regime', 'Regime luminansi latar tampak', 'regime', '—', 'Scotopic <0.005; mesopic 0.005–5; photopic >5 cd/m². Diagnostik, bukan riwayat adaptasi.', DECIMAL),
    ('achromatic_extrapolation', 'Extrapolasi achromatic senja', 'achromatic_extrapolation', 'boolean', 'Mesopic/photopic: aproksimasi threshold luminansi tanpa model warna hilal/langit.', DECIMAL),
    ('B_floor', 'Batas floor background', 'B_floor', 'cd/m²', '1×10⁻⁵; kebijakan membekukan increment pada area aktual.', SCIENTIFIC),
    ('B_eval', 'Latar untuk evaluasi koefisien', 'B_eval', 'cd/m²', 'max(B_app, B_floor); C_th tetap memakai latar aktual.', SCIENTIFIC),
    ('floor', 'Floor aktif', 'floor_applied', 'boolean', 'Perluasan implementasi; berbeda dari cutoff M₀ literal di paper.', DECIMAL),
    ('R', 'Parameter Ricco', 'R(B_eval)', 'sr', 'Koefisien persamaan Crumey sesuai kontrak threshold yang sama.', SCIENTIFIC),
    ('C_inf', 'Kontras batas target luas', 'C∞(B_eval)', '—', 'Kontras asimtotik sebelum multiplier φ.', SCIENTIFIC),
    ('q', 'Eksponen penyambung', 'q(B_eval)', '—', 'Eksponen generalized mean pada threshold.', DECIMAL),
    ('ricco_area_sr', 'Luas Ricco', 'A_R = R/C∞', 'sr', 'Ukuran karakteristik kurva, bukan pengganti luas sabit.', SCIENTIFIC),
    ('delta_app', 'Ambang increment pada mata', 'ΔB_app,th', 'cd/m²', 'φ B_eval [(R/A_app)^q + C∞^q]^(1/q).', SCIENTIFIC),
    ('delta_sky', 'Ambang increment ekuivalen di langit', 'ΔB_sky,th', 'cd/m²', 'ΔB_app,th / g; +inf untuk optik terblokir.', SCIENTIFIC),
    ('C_th', 'Ambang kontras', 'C_th', '—', 'ΔB_sky,th / B_sky = ΔB_app,th / B_app jika g > 0.', SCIENTIFIC),
    '06 | PERBANDINGAN DAN KEPUTUSAN MODEL',
    ('C_obj', 'Kontras excess hilal', 'C_obj', '—', 'L / B_sky; sama untuk NE/teleskop pada waktu yang sama jika optik meneruskan cahaya.', SCIENTIFIC),
    ('ratio', 'Rasio terhadap ambang', 'C_obj/C_th', '—', '> 1 berarti melampaui ambang; = 1 belum diklasifikasikan terlihat.', DECIMAL),
    ('log_margin', 'Margin logaritmik', 'log₁₀(C_obj/C_th)', '—', 'Δm / 2.5.', DECIMAL),
    ('margin', 'Margin visibilitas', 'Δm', 'mag', '2.5 log₁₀(C_obj/C_th). Positif: melampaui ambang model.', DECIMAL),
    ('gain', 'Selisih threshold NE/teleskop', 'G', 'mag', '2.5 log₁₀(C_th,NE/C_th,tel), waktu sama; mencakup perubahan F jika dua nilai berbeda.', DECIMAL),
    ('field_factor_comparison', 'Kontrak perbandingan residual F', 'F_comparison', '—', 'shared_residual_factor: base F sama; independent_residual_factors: dua F eksplisit berbeda.', DECIMAL),
    ('status', 'Keputusan visibilitas', 'Δm > 0', '—', 'Prediksi visual model; bukan label pengamatan atau probabilitas.', DECIMAL),
]


def _write_rows(ws, values, rows=CHAIN, start=6):
    row = start
    for entry in rows:
        if isinstance(entry, str):
            _bar(ws, row, entry)
        else:
            key, label, symbol, unit, note, fmt = entry
            for col, val in enumerate([label, symbol, unit, *[v.get(key) for v in values], note], 1):
                fill = 'F4FAFB' if col in (4, 5) else 'F8F5FC' if col in (6, 7) else ('F8FAFC' if row % 2 == 0 else None)
                if key == 'status' and col in range(4, 8):
                    fill = GREEN if val == 'TERLIHAT (MODEL)' else RED if val and 'TIDAK' in val else GREY
                _cell(ws, row, col, val, fmt, fill, key in ('status', 'margin'))
            ws.row_dimensions[row].height = 46 if len(note) > 88 else 34
            if key == 'margin':
                area = f'D{row}:G{row}'
                for op, color in (('greaterThan', GREEN), ('lessThanOrEqual', RED)):
                    ws.conditional_formatting.add(area, CellIsRule(operator=op, formula=['0'], fill=PatternFill('solid', fgColor=color)))
        row += 1
    ws.print_area = f'A1:H{row - 1}'
    return row


def _input_sheet(calc, ws):
    _page(ws, 'INPUT DAN KONFIGURASI', 'Nilai yang digunakan saat perhitungan; satuan dipisahkan dari nilai.', [34, 23, 20, 38, 84])
    for c, title in enumerate(['Parameter', 'Simbol', 'Satuan', 'Nilai', 'Keterangan'], 1):
        _cell(ws, 4, c, title, fill=PALE, bold=True)
    h, tel = calc.hasil, calc.hasil.get('tel_params') or {}
    scan = h.get('scan_params') or {}
    rows = [
        'LOKASI DAN PENANGGALAN',
        ('Lokasi', 'nama', '—', calc.nama_tempat, 'Satu lokasi pengamatan.'),
        ('Lintang', 'latitude', 'deg', calc.lintang, 'Positif utara.'),
        ('Bujur', 'longitude', 'deg', calc.bujur, 'Positif timur.'),
        ('Elevasi', 'H', 'm', calc.elevasi, 'Ketinggian lokasi di atas muka laut.'),
        ('Zona waktu', 'timezone', '—', calc.timezone_str, 'Seluruh waktu lokal di workbook menggunakan zona ini.'),
        ('Zona acuan H+0', 'timezone_H0', '—', h.get('h0_reference_timezone', 'Asia/Jakarta'), 'Tanggal ijtima WIB adalah H+0 untuk semua lokasi; batas hari 00:00 WIB.'),
        ('Bulan Hijri', 'bulan', '—', calc.bulan_hijri, ''),
        ('Tahun Hijri', 'tahun', '—', calc.tahun_hijri, ''),
        ('Offset hari', 'offset', 'hari', calc.delta_day_offset, 'Offset tanggal; bukan jam pengamatan.'),
        ('Tanggal pengamatan', 'date', 'lokal', h.get('tanggal_pengamatan'), ''),
        ('Ijtima UTC', 't_ijtima', 'UTC', h.get('ijtima_utc'), 'Nilai waktu UTC; sel berbeda dengan waktu lokal.'),
        ('Ijtima lokal', 't_ijtima', 'lokal', h.get('ijtima_local'), ''),
        ('Sunset estimasi untuk atmosfer', 't_est', 'lokal', h.get('sunset_est_local'), 'Timestamp pengambilan atmosfer untuk perhitungan sunset apparent.'),
        ('Sunset apparent', 't_sunset', 'lokal', h.get('sunset_local'), 'Sunset dengan refraksi menggunakan T/P yang tersimpan.'),
        ('Sunset apparent UTC', 't_sunset', 'UTC', h.get('sunset_utc'), ''),
        'ATMOSFER DAN PENCARIAN WAKTU',
        ('Sumber atmosfer', 'source', '—', calc.SUMBER_ATMOSFER_LABEL.get(calc.sumber_atmosfer), 'Provenance mentah juga disimpan pada file .xlsx.atmosphere.json.'),
        ('RH awal sunset', 'RH_raw', '%', h.get('rh_raw'), 'Sebelum koreksi bias; timestamp sunset estimasi.'),
        ('Suhu awal sunset', 'T_raw', '°C', h.get('temperature_raw'), 'Sebelum koreksi bias; timestamp sunset estimasi.'),
        ('Bias RH dikonfigurasi', 'bias_RH', 'poin %', h.get('bias_rh'), 'Bias = reanalisis − observasi; tidak diterapkan pada sumber manual.'),
        ('Bias suhu dikonfigurasi', 'bias_T', '°C', h.get('bias_t'), 'Bias = reanalisis − observasi; tidak diterapkan pada sumber manual.'),
        ('Mode waktu', 'mode', '—', h.get('mode'), 'sunset atau optimal; berbeda dari mode lokasi tunggal/multi.'),
        ('Interval scan', 'Δt_scan', 'menit', scan.get('interval_menit'), 'Digunakan hanya pada mode optimal.'),
        ('Delay awal', 'delay', 'menit', scan.get('start_delay_menit'), 'Scan dimulai setelah sunset sesuai delay ini.'),
        ('Batas altitude Bulan', 'h_min', 'deg', scan.get('min_moon_alt'), 'Berlaku pada scan dan refinement.'),
        ('Langkah refinement', 'Δt_refine', 'detik', scan.get('refinement_seconds'), 'Pencarian lebih rapat di sekitar kandidat optimum.'),
        'PENGAMAT DAN TELESKOP',
        ('Teleskop diaktifkan', 'use_telescope', 'boolean', h.get('use_telescope', True), 'Jika False, kolom hasil teleskop kosong; bukan margin nol.'),
        ('Residual field factor NE', 'F_NE', '—', h.get('F_naked'), 'Asumsi referensi, belum dikalibrasi khusus hilal; efek fisik eksplisit tidak dihitung ulang.'),
        *[(label, symbol, unit, tel.get(key), note) for key, label, symbol, unit, note in (
            ('aperture', 'Diameter aperture', 'D_ap', 'mm', 'Berbeda dari luas D_luas pada Kastner.'),
            ('magnification', 'Pembesaran', 'M', '×', ''),
            ('central_obstruction', 'Obstruksi pusat', 'D_s', 'mm', ''),
            ('transmission', 'Transmisi per permukaan', 't₁', 'fraksi', 'Transmisi total τ = t₁ⁿ.'),
            ('n_surfaces', 'Jumlah permukaan optik', 'n', 'permukaan', ''),
            ('observer_age', 'Usia pengamat', 'age', 'tahun', 'Hanya untuk fallback pupil; bukan koreksi sensitivitas umur.'),
            ('pupil_diameter_mm', 'Override pupil mata', 'p_override', 'mm', 'Kosong: gunakan estimasi pupil gelap berdasarkan usia, bukan pengukuran pupil senja.'),
            ('field_factor', 'Residual field factor teleskop', 'F_tel', '—', 'Default mengikuti F_NE; multiplier monokular √2 diterapkan secara terpisah.'),
        )],
        'KONSTANTA DAN ASUMSI AKTIF',
        ('Zero point V', 'Z_V', 'lux', Z_V, 'Fluks V = 0; sama untuk seluruh identitas fotometri.'),
        ('Konversi S10 ke nL', 'S10_to_nL', 'nL/S10', S10_TO_NL, 'Berasal dari Z_V dan satu derajat persegi.'),
        ('Konversi nL ke SI', 'nL_to_cdm2', 'cd/m² per nL', nL_to_cdm2(1), ''),
        ('Floor background', 'B_floor', 'cd/m²', B_FLOOR, 'Increment dibekukan di floor untuk area aktual; kebijakan perluasan implementasi.'),
        ('Faktor monokular teleskop', 'F_T', '—', math.sqrt(2), 'Sensitivitas monokular menaikkan threshold.'),
        ('Faktor tambahan', 'F_M', '—', 1, 'Seeing, PSF, glare dan medan pandang belum dimodelkan eksplisit.'),
    ]
    for row, entry in enumerate(rows, 5):
        if isinstance(entry, str):
            _bar(ws, row, entry, 5)
        else:
            for col, value in enumerate(entry, 1):
                _cell(ws, row, col, value, SCIENTIFIC if entry[1] in ('Z_V', 'B_floor', 'nL_to_cdm2') else DECIMAL,
                      PALE if row % 2 == 0 else None)
            ws.row_dimensions[row].height = 34
    ws.freeze_panes = 'D5'
    ws.print_title_rows = '1:4'
    ws.print_area = f'A1:E{ws.max_row}'


def _atmosphere_sheet(calc, ws, scenes):
    _page(ws, 'RINCIAN ATMOSFER — SCHAEFER', 'Satu snapshot sunset, optimal mata telanjang, dan optimal teleskop. Rincian ini direkam saat model dihitung.', [34, 21, 18, 25, 25, 25, 76])
    for col, title in enumerate(['Parameter', 'Simbol', 'Satuan', 'Sunset', 'Optimal mata telanjang', 'Optimal teleskop', 'Rumus / keterangan'], 1):
        _cell(ws, 4, col, title, fill=PALE, bold=True)
    entries = [
        ('waktu_local', 'Waktu evaluasi lokal', 't', 'lokal', 'Sunset menggunakan sunset_local.'),
        ('rh', 'Kelembapan', 'RH', '%', 'Input aktual Schaefer.'),
        ('temperature', 'Suhu', 'T', '°C', 'Input aktual Schaefer.'),
        ('pressure', 'Tekanan', 'P', 'hPa', 'Input refraksi; tidak masuk eksplisit ke rumus ekstingsi.'),
        *[(key, label, symbol, unit, note) for key, label, symbol, unit, note in (
            ('Z', 'Jarak zenit objek untuk Schaefer', 'Z', 'deg', '90 − max(h_moon, 0); sama dengan input actual Schaefer.'),
            ('Z_sun', 'Jarak zenit Matahari', 'Z_sun', 'deg', '90 − h_sun.'),
            ('separation_sky_deg', 'Separasi untuk rumus langit', 'R_sun', 'deg', 'Dari altitude dan selisih azimuth input Schaefer; dapat berbeda sedikit dari E astrometrik.'),
            ('X_G', 'Airmass gas', 'X_G', '—', 'Dipakai untuk Rayleigh dan uap air.'),
            ('X_A', 'Airmass aerosol', 'X_A', '—', 'Dipakai untuk aerosol.'),
            ('X_O', 'Airmass ozon', 'X_O', '—', 'Geometri lapisan ozon.'),
            ('X_sky', 'Airmass umum latar', 'X', '—', 'Rumus latar memakai X, bukan DM[2].'),
            ('X_sun', 'Airmass Matahari', 'X_sun', '—', 'Model menetapkan 40 ketika Matahari di bawah horizon.'),
        )],
    ]
    row = 5
    flattened = []
    for scene in scenes:
        scene = scene or {}
        atmosphere = (scene.get('model_trace') or {}).get('atmosphere') or {}
        flattened.append({**scene, 'waktu_local': scene.get('waktu_local') or scene.get('sunset_local'),
                          **(atmosphere.get('diagnostics') or {})})
    def write(entry, values):
        nonlocal row
        key, label, symbol, unit, note = entry
        for col, value in enumerate([label, symbol, unit, *values, note], 1):
            _cell(ws, row, col, value, SCIENTIFIC if unit in ('nL', 'fraksi', 'mag/airmass') else DECIMAL,
                  PALE if row % 2 == 0 else None)
        ws.row_dimensions[row].height = 34
        row += 1
    for entry in entries:
        write(entry, [v.get(entry[0]) for v in flattened])
    for band, index in zip(('U', 'B', 'V', 'R', 'I'), range(5)):
        _bar(ws, row, f'BAND {band} — KOEFISIEN DAN EKSTINGSI', 7, TEAL if band == 'V' else NAVY)
        row += 1
        for key, label in (('k_R', 'Rayleigh'), ('k_A', 'Aerosol'), ('k_O', 'Ozon'), ('k_W', 'Uap air')):
            write((key, label, key, 'mag/airmass', 'Komponen yang direkam oleh perhitungan per band.'),
                  [(v.get('bands', {}).get(band) or {}).get(key) for v in flattened])
        for key, label, unit in (('K', 'Koefisien total', 'mag/airmass'), ('DM', 'Ekstingsi total sepanjang LOS', 'mag')):
            vals = []
            for scene in scenes:
                array = (((scene or {}).get('model_trace') or {}).get('atmosphere') or {}).get(key) or []
                vals.append(array[index] if len(array) > index else None)
            write((key, label, f'{key}[{band}]', unit, 'K = Σ k_i; DM = Σ k_i X_i.'), vals)
    _bar(ws, row, 'BAND V — PEMBENTUKAN LATAR', 7, TEAL)
    row += 1
    for key, label, unit, note in (
        ('T_sky', 'Transmisi dalam rumus latar', 'fraksi', '10^(−0.4 k_v X); berbeda dari T_V sumber.'),
        ('T_sun', 'Transmisi lintasan Matahari', 'fraksi', '10^(−0.4 k_v X_sun).'),
        ('B_night_nL', 'Komponen night sky', 'nL', 'BN, termasuk atenuasi dalam rumus Schaefer.'),
        ('B_twilight_nL', 'Kandidat twilight', 'nL', 'BT; salah satu kandidat komponen latar.'),
        ('B_daylight_nL', 'Kandidat daylight', 'nL', 'BD; salah satu kandidat komponen latar.'),
    ):
        write((key, label, key, unit, note), [(v.get('bands', {}).get('V') or {}).get(key) for v in flattened])
    write(('B_sky', 'Latar total', 'B_sky', 'nL', 'BN + min(BT, BD); tanpa moonlight; tanpa perkalian T_V tambahan.'),
          [(scene or {}).get('sky_brightness_nl') for scene in scenes])
    ws.freeze_panes = 'D5'
    ws.print_title_rows = '1:4'
    ws.print_area = f'A1:G{row - 1}'


def _timestep_sheet(calc, ws):
    """One filterable numeric table; shared, NE and telescope blocks are separate."""
    shared = [
        ('No', None), ('Waktu Lokal', 'time'), ('Moon Alt (°)', 'moon_alt'), ('Sun Alt (°)', 'sun_alt'),
        ('Elongasi (°)', 'elongation'), ('Sudut fase (°)', 'phase_angle'), ('Semidiameter Bulan (deg)', 'moon_semidiameter'),
        ('Jarak Bulan (km)', 'moon_distance_km'), ('RH (%)', 'rh'), ('T (°C)', 'temperature'), ('P (hPa)', 'pressure'),
        ('k_v (mag/airmass)', 'k_v'), ('extinction_mag_v (mag)', 'extinction_mag_v'), ('transmission_v', 'transmission_v'),
        ('M_v intrinsik (mag)', 'M_v'), ('L* (S10)', 'L_star_s10'), ('D_luas (deg²)', 'area_deg2'), ('A (sr)', 'area_sr'),
    ]
    ne_cols = [('NE | L (nL)', 'L_app_nL'), ('NE | B (nL)', 'B_app_nL'), ('NE | C_obj', 'C_obj'),
               ('NE | C_th', 'C_th'), ('NE | ΔB_th (cd/m²)', 'delta_app'), ('Δm Naked Eye', 'margin'),
               ('NE | Regime', 'regime'), ('NE | Extrapolasi achromatic', 'achromatic_extrapolation'), ('NE | Keputusan', 'status')]
    tel_cols = [('TEL | g', 'g'), ('TEL | A_app (sr)', 'area_app_sr'), ('TEL | L_app (nL)', 'L_app_nL'),
                ('TEL | B_app (nL)', 'B_app_nL'), ('TEL | C_obj', 'C_obj'), ('TEL | C_th', 'C_th'),
                ('TEL | ΔB_app,th (cd/m²)', 'delta_app'), ('Margin Teleskop', 'margin'),
                ('TEL | Regime', 'regime'), ('TEL | Extrapolasi achromatic', 'achromatic_extrapolation'),
                ('TEL | Keputusan', 'status'), ('Selisih Threshold (mag)', 'gain'), ('TEL | Perbandingan F', 'field_factor_comparison')]
    cols = shared + ne_cols + tel_cols + [('Status geometri', None)]
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TEAL
    ws.freeze_panes = 'C2'
    ws.print_title_rows = '1:1'
    ws.row_dimensions[1].height = 46
    for col, (label, key) in enumerate(cols, 1):
        color = NAVY if col <= len(shared) else TEAL if col <= len(shared) + len(ne_cols) else PURPLE
        cell = _cell(ws, 1, col, label, fill=color, bold=True)
        cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.comment = Comment('Blok input bersama / mata telanjang / teleskop. Angka finite disimpan sebagai angka; inf/-inf sebagai teks.', 'Model hilal')
        ws.column_dimensions[get_column_letter(col)].width = 22 if col == 2 else 18
    for start, end in ((3, len(shared)), (len(shared) + 1, len(shared) + len(ne_cols)),
                       (len(shared) + len(ne_cols) + 1, len(cols) - 1)):
        ws.column_dimensions.group(get_column_letter(start), get_column_letter(end), outline_level=1, hidden=False)
    results = calc.hasil.get('all_timestep_results') or []
    for row, scene in enumerate(results, 2):
        ne, tel = _scene_values(calc, scene), _scene_values(calc, scene, True)
        # Keep geometry/weather even for a rejected scene; no synthetic -99 margin.
        base = {**scene, **ne, 'time': scene.get('waktu_local')}
        values = [row - 1] + [base.get(key) for _, key in shared[1:]]
        values += [ne.get(key) for _, key in ne_cols] + [tel.get(key) for _, key in tel_cols]
        values += ['VALID' if scene.get('valid') else 'DI BAWAH HORIZON / TIDAK DIEVALUASI']
        for col, (value, (_, key)) in enumerate(zip(values, cols), 1):
            fill = GREEN if value == 'TERLIHAT (MODEL)' else RED if value == 'TIDAK TERLIHAT (MODEL)' else None
            _cell(ws, row, col, value, SCIENTIFIC if key in ('g', 'area_app_sr', 'L_star_s10', 'area_deg2', 'area_sr',
                  'transmission_v', 'C_obj', 'C_th', 'delta_app', 'L_app_nL', 'B_app_nL') else DECIMAL, fill)
    if results:
        ref = f'A1:{get_column_letter(len(cols))}{len(results) + 1}'
        table = Table(displayName='DataWaktuHilal', ref=ref)
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(table)
    else:
        _cell(ws, 2, 1, 'Mode sunset: hasil lengkap ada pada Ringkasan dan Rantai Model.')


def _chart(calc, summary, wb, anchor):
    results = calc.hasil.get('all_timestep_results') or []
    sunset = calc.hasil.get('sunset_local')
    if not results or not sunset:
        return
    data = wb.create_sheet('_Data Grafik')
    data.append(['Menit setelah sunset', NE_LABEL, TEL_LABEL, 'Ambang Δm = 0'])
    # Include the final refined optima: they may not lie on the scan grid.
    samples = {r['waktu_local']: r for r in results if r.get('valid') and r.get('waktu_local')}
    for key in ('optimal_result_ne', 'optimal_result_tel'):
        scene = calc.hasil.get(key)
        if scene and scene.get('valid'):
            samples[scene['waktu_local']] = scene
    if not samples:
        wb.remove(data)
        return
    for time, scene in sorted(samples.items()):
        finite = lambda x: x if isinstance(x, (float, int)) and math.isfinite(x) else None
        data.append([(time - sunset).total_seconds() / 60, finite(scene.get('delta_m_ne')),
                     finite(scene.get('delta_m_tel')) if calc.hasil.get('use_telescope', True) else None, 0])
    chart = ScatterChart()
    chart.title = 'Margin visibilitas sepanjang waktu'
    chart.x_axis.title, chart.y_axis.title = 'Menit setelah sunset', 'Δm (mag)'
    chart.style, chart.height, chart.width = 13, 12, 30
    chart.display_blanks = 'gap'
    chart.visible_cells_only = False
    x = Reference(data, min_col=1, min_row=2, max_row=data.max_row)
    for col, color in ((2, TEAL), (3, PURPLE), (4, '7C8B99')):
        if col == 3 and not calc.hasil.get('use_telescope', True):
            continue
        series = Series(Reference(data, min_col=col, min_row=1, max_row=data.max_row), x, title_from_data=True)
        series.graphicalProperties.line.solidFill = color
        series.graphicalProperties.line.width = 24000
        if col == 4:
            series.graphicalProperties.line.prstDash = 'dash'
        chart.series.append(series)
    summary.add_chart(chart, f'A{anchor}')
    data.sheet_state = 'hidden'


def _info_sheet(calc, ws):
    _page(ws, 'PANDUAN MEMBACA HASIL', 'Prediksi model visual; perbaikan rumus dan tes numerik terpisah dari validasi empiris.', [34, 112])
    sections = [
        ('Mulai dari Ringkasan', 'Kolom D/E: mata telanjang/teleskop saat sunset. Kolom F/G: optimum masing-masing. Jika waktu optimal berbeda, kondisi langitnya juga berbeda.'),
        ('Telusuri Rantai Model', 'Geometri → magnitudo dan L* → ekstingsi dan L → latar Schaefer → optik → threshold Crumey → rasio dan Δm → keputusan.'),
        ('Periksa Input & Konfigurasi', 'Lokasi, waktu, sumber atmosfer, bias, konfigurasi teleskop dan faktor pengamat yang benar-benar digunakan.'),
        ('Periksa Atmosfer', 'Airmass, komponen koefisien ekstingsi U/B/V/R/I dan kandidat night/twilight/daylight band V. Transmisi latar dan transmisi sumber berbeda.'),
        ('Gunakan Timestep Data', 'Tabel dapat difilter dan kolom dikelompokkan. Input bersama, hasil NE, dan hasil TEL berada dalam blok kolom terpisah.'),
        ('Angka dan satuan', 'Finite disimpan sebagai angka dengan presisi sumber; tampilan dibulatkan melalui format sel saja. Sel waktu adalah wall time lokal/UTC sesuai label.'),
        ('Nilai kosong', 'Parameter tidak berlaku, diagnostik belum direkam, atau optimum tidak tersedia. Kosong tidak sama dengan nol.'),
        ('inf / -inf', 'Teks karena Excel tidak menyimpan infinity. Sumber nol memberi margin −inf; optik terblokir memberi C_th +inf. Grafik mengosongkan nilai nonfinite.'),
        ('Batas keputusan', 'Δm > 0 melampaui threshold; Δm = 0 belum terlihat. Bulan di bawah horizon diberi status geometris tersendiri. Ini bukan peluang deteksi.'),
        ('Luminansi hilal', 'L adalah increment cahaya hilal di atas latar, bukan L+B. Karena itu C_obj = L/B, bukan (L−B)/B.'),
        ('D dan Δm', 'D_luas [deg²] di Kastner berbeda dari D_ap [mm] di teleskop. A_V = DM[2] [mag] adalah ekstingsi total; Δm [mag] adalah margin visibilitas.'),
        ('Model optik', 'Crumey extended source dengan g annular dan A_app=M²A. Faktor historis Fa/Fm/Fr bukan multiplier aktif; seeing tidak diterapkan.'),
        ('Kurva dan floor', 'auto memakai combined. B_eval=max(B_app,10⁻⁵). Floor membekukan increment untuk area aktual; merupakan perluasan yang dinyatakan, bukan cutoff M₀ literal paper.'),
        ('Residual F dan selisih threshold', 'Base F default sama untuk NE/teleskop, belum dikalibrasi khusus hilal. Atmosfer, throughput, magnifikasi dan FT/FM dihitung terpisah. Dua F eksplisit berbeda ikut mengubah selisih threshold.'),
        ('Regime dan extrapolasi achromatic', 'Regime memakai latar aktual/apparent, bukan floor dan bukan riwayat adaptasi. Flag mesopic/photopic menyatakan extrapolasi luminansi achromatic tanpa validasi persepsi warna hilal senja.'),
        ('Asumsi hilal', 'Luas sabit dan luminansi rata-rata mengabaikan bentuk tipis melengkung dan distribusi luminansi. Phase law dekat fase 180° dan koreksi jarak fluks masih memiliki keterbatasan.'),
        ('Pupil dan adaptasi', 'Pupil fallback dari usia adalah estimasi pupil gelap, bukan pupil senja terukur. Adaptasi, warna senja, glare, PSF dan luas yang muat di medan belum tervalidasi untuk hilal.'),
        ('Durasi window', 'Durasi model adalah panjang streak scan × interval, bukan durasi pengamatan yang diukur. Waktu tepi dibatasi grid scan; refinement hanya mencari puncak.'),
        ('Status validasi empiris', '278 label BMKG dikonfirmasi sebagai CCD/citra digital. Konfigurasi dan waktu aktual tidak tersedia; data tersebut tidak mengkalibrasi sensitivitas visual Crumey.'),
        ('Snapshot perhitungan', 'Parameter perantara direkam di model_trace saat perhitungan. Export tidak memanggil API cuaca atau memilih konfigurasi baru. Diagnostik hasil lama yang belum tersimpan dibiarkan kosong.'),
        ('Ephemeris', 'Skyfield + DE440s; semidiameter dari jarak astrometrik toposentrik. Sudut fase dan pendekatan luas elongasi dijelaskan terpisah.'),
        ('Model latar', 'Schaefer (1993), adaptasi tanpa moonlight karena Bulan adalah target. Atmosfer sudah masuk di rumus latar; tidak ditransmisikan ulang dengan T_V sumber.'),
        ('Model fotometri', 'Kastner / phase law yang dipertahankan di repo; memakai konversi S10 dan zero point V bersama.'),
        ('Model ambang', 'Crumey (2014), Human contrast threshold and astronomical visibility, DOI 10.1093/mnras/stu992.'),
        ('Waktu ekspor UTC', datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')),
    ]
    for row, (label, value) in enumerate(sections, 4):
        _cell(ws, row, 1, label, fill=PALE, bold=True)
        _cell(ws, row, 2, value)
        ws.row_dimensions[row].height = 42
    ws.freeze_panes = 'B4'
    ws.print_area = f'A1:B{ws.max_row}'


def write_single_workbook(calc, filepath):
    if not calc.hasil or 'delta_m_ne' not in calc.hasil:
        raise ValueError('Jalankan perhitungan sebelum mengekspor Excel.')
    h = calc.hasil
    sunset, opt_ne, opt_tel = h, h.get('optimal_result_ne'), h.get('optimal_result_tel')
    if h.get('mode') != 'optimal':
        opt_ne = opt_tel = None
    if not h.get('use_telescope', True):
        opt_tel = None
    values = [_scene_values(calc, sunset), _scene_values(calc, sunset, True),
              _scene_values(calc, opt_ne), _scene_values(calc, opt_tel, True)]
    wb = Workbook()
    wb.properties.title = 'Visibilitas hilal — satu lokasi'
    wb.properties.subject = 'Rantai Schaefer → Kastner → optik → Crumey'
    wb.properties.creator = 'HilalVisibilityCalculator'
    summary = wb.active
    summary.title = 'Ringkasan'
    subtitle = f'{calc.nama_tempat} | {calc.bulan_hijri}/{calc.tahun_hijri} H | mode {h.get("mode", "sunset")} | zona waktu {calc.timezone_str}'
    _matrix(summary, 'VISIBILITAS HILAL — RINGKASAN', subtitle)
    selected = ('status', 'time', 'margin', 'ratio', 'moon_alt', 'sun_alt', 'elongation', 'phase_angle',
                'L_app_nL', 'B_app_nL', 'C_obj', 'C_th', 'gain')
    row = _write_rows(summary, values, [next(x for x in CHAIN if isinstance(x, tuple) and x[0] == key) for key in selected])
    _bar(summary, row + 1, 'WINDOW MODEL — MODE OPTIMAL')
    windows = [
        ('visibility_start_', 'Awal terlihat pada grid', 'lokal'), ('visibility_end_', 'Akhir terlihat pada grid', 'lokal'),
        ('best_window_start_', 'Awal streak terpanjang', 'lokal'), ('best_window_end_', 'Akhir streak terpanjang', 'lokal'),
        ('visibility_duration_', 'Durasi streak model', 'menit'),
    ]
    for r, (key, label, unit) in enumerate(windows, row + 2):
        vals = [label, key.rstrip('_'), unit, None, None,
                h.get(key + 'ne') if opt_ne else None, h.get(key + 'tel') if opt_tel else None,
                'Estimasi grid scan; bukan durasi pengamatan aktual.']
        for col, val in enumerate(vals, 1):
            _cell(summary, r, col, val)
        summary.row_dimensions[r].height = 32
    note_row = row + 8
    _bar(summary, note_row, 'CARA MEMBACA HASIL', color=TEAL)
    summary.merge_cells(start_row=note_row + 1, start_column=1, end_row=note_row + 1, end_column=8)
    _cell(summary, note_row + 1, 1,
          'Δm > 0: melampaui ambang model visual. Sel kosong: tidak berlaku / optimum tidak tersedia / belum direkam. '
          'Ini bukan validasi pengamatan CCD. Lihat Rantai Model untuk semua besaran dan Info Program untuk batas asumsi.')
    summary.row_dimensions[note_row + 1].height = 38
    for col, name in enumerate(('Input & Konfigurasi', 'Rantai Model', 'Atmosfer', 'Timestep Data', 'Info Program'), 1):
        cell = _cell(summary, note_row + 2, col, name, bold=True)
        cell.hyperlink = f"#'{name}'!A1"
        cell.font = Font(name='Calibri', size=10, color=TEAL, underline='single')
    _input_sheet(calc, wb.create_sheet('Input & Konfigurasi'))
    chain = wb.create_sheet('Rantai Model')
    _matrix(chain, 'RANTAI PEMODELAN — DARI INPUT KE KEPUTUSAN', subtitle)
    _write_rows(chain, values)
    _atmosphere_sheet(calc, wb.create_sheet('Atmosfer'), (sunset, opt_ne, opt_tel))
    _timestep_sheet(calc, wb.create_sheet('Timestep Data'))
    _info_sheet(calc, wb.create_sheet('Info Program'))
    _chart(calc, summary, wb, note_row + 4)
    if summary._charts:
        summary.print_area = f'A1:H{note_row + 29}'
    output = Path(filepath)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    wb.close()
    return str(output)
