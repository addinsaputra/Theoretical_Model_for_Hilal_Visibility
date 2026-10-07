#!/usr/bin/env python3
"""
======================================================================
ANALISIS DIAGNOSTIK MODEL CRUMEY (2014)
Tahap 4.5: Analisis Sensitivitas dan Ketidakpastian
======================================================================
Versi dimodifikasi untuk konsistensi dengan arsitektur Bab 4:
  - F_NAKED = F_TEL = 2.0 (referensi sensitivitas, belum dikalibrasi visual)
  - TEL_AGE = 22.0 (konsisten dengan batch run)
  - sun_alt dari data CSV (bukan estimasi)
  - Input: data_hilal_era5.csv, data_hilal_merra2.csv
Label CCD BMKG hanya mendukung perbandingan deskriptif lintas metode;
label tersebut tidak mengkalibrasi threshold penglihatan manusia.

CARA PAKAI:
  1. Gunakan checkout proyek beserta dependensinya
  2. Pastikan file CSV ada di outputs/
  3. Jalankan: python scripts/analyze_crumey.py
======================================================================
"""

import math
import os
import sys
import warnings
warnings.filterwarnings('ignore')

import pandas as pd
import numpy as np
from typing import List, Dict, Tuple, Optional


from hilal_visibility.paths import OUTPUT_DIR as PROJECT_OUTPUT_DIR
from hilal_visibility.models.schaefer import hitung_sky_brightness
from hilal_visibility.models.kastner import hitung_luminansi_intrinsik, terapkan_transmisi_atmosfer
from hilal_visibility.models.crumey import (
    nL_to_cdm2, cdm2_to_nL, arcmin2_to_sr,
    contrast_threshold, crescent_area_arcmin2,
    visibility_margin_mag, telescopic_extended_threshold,
)
from hilal_visibility.models.telescope import TelescopeVisibilityModel
from hilal_visibility.ephemeris import moon_semidiameter_time_local, set_location
from hilal_visibility.calculator import tentukan_timezone_indonesia
from hilal_visibility.console import configure_console_encoding
from pytz import timezone


# ===================================================================
# KONFIGURASI
# ===================================================================

OUTPUT_DIR = str(PROJECT_OUTPUT_DIR)

INPUT_ERA5 = os.path.join(OUTPUT_DIR, 'data_hilal_era5.csv')
INPUT_MERRA2 = os.path.join(OUTPUT_DIR, 'data_hilal_merra2.csv')

RH_DELTAS = [-30, -25, -20, -15, -10, -5, 0, +5, +10]

# Telescope parameters (konsisten dengan batch run)
TEL_APERTURE = 100.0     # mm
TEL_MAG = 50.0
TEL_TRANS = 0.95
TEL_NSURFACES = 6
TEL_OBSTRUCTION = 0.0
TEL_AGE = 22.0           # <-- FIX: konsisten dengan batch run

# Reference factors for sensitivity calculations, not empirically calibrated
# visual thresholds. BMKG gallery labels are CCD detections, not human vision.
F_NAKED = 2.0
F_TEL = 2.0

PROVENANCE_COLUMNS = (
    'Observation_Method', 'Observation_Source', 'Comparison_Scope',
    'Actual_Telescope_Config_Available', 'Actual_Observation_Time_Available',
)


def _observation_provenance(row):
    """Carry input metadata into derived sensitivity/report artifacts."""
    result = {}
    for column in PROVENANCE_COLUMNS:
        default = 'unknown' if column in PROVENANCE_COLUMNS[:3] else None
        value = row.get(column, default)
        result[column] = default if pd.isna(value) else value
    return result

# ===================================================================
# DATA LOADING
# ===================================================================

def load_observation_data(filepath: str) -> pd.DataFrame:
    """Load usable model inputs, retain observation provenance and exclusions.

    Missing method metadata stays unknown; model/label agreement remains a
    descriptive comparison and cannot establish a visual detection threshold.
    Failed atmosphere rows never reach ephemeris reconstruction or sensitivity.
    """
    df_raw = pd.read_csv(filepath, float_precision='round_trip')
    reasons = pd.Series('', index=df_raw.index, dtype='object')

    def exclude(mask, reason):
        reasons.loc[mask & reasons.eq('')] = reason

    if 'Status' in df_raw:
        exclude(~df_raw['Status'].fillna('').astype(str).str.strip().str.lower().eq('valid'),
                'status_invalid')
    phase_column = 'Phase_Angle_BT' if 'Phase_Angle_BT' in df_raw else 'Phase_Angle'
    numeric_inputs = (
        'No', 'Lat', 'Lon', 'Elv', phase_column, 'sun_alt_BT', 'Moon_Alt_BT',
        'Elongasi_BT', 'Sky_Bright_BT', 'Lum_Hilal_BT', 'kV_BT', 'RH_BT', 'T_BT',
    )
    for column in numeric_inputs:
        if column not in df_raw:
            raise ValueError(f'CSV diagnostic memerlukan kolom {column}.')
        df_raw[column] = pd.to_numeric(df_raw[column], errors='coerce')
        exclude(~np.isfinite(df_raw[column]), f'nonfinite_input:{column}')
    for column, valid in (
        ('No', (df_raw['No'] >= 0) & (df_raw['No'] % 1 == 0)),
        ('Lat', df_raw['Lat'].between(-90, 90)),
        ('Lon', df_raw['Lon'].between(-180, 180)),
        (phase_column, df_raw[phase_column].between(0, 180)),
        ('sun_alt_BT', df_raw['sun_alt_BT'].between(-90, 90)),
        ('Moon_Alt_BT', (df_raw['Moon_Alt_BT'] > 0) & (df_raw['Moon_Alt_BT'] <= 90)),
        ('Elongasi_BT', df_raw['Elongasi_BT'].between(0, 180)),
        ('Sky_Bright_BT', df_raw['Sky_Bright_BT'] > 0),
        ('Lum_Hilal_BT', df_raw['Lum_Hilal_BT'] >= 0),
        ('RH_BT', (df_raw['RH_BT'] >= 0) & (df_raw['RH_BT'] < 100)),
    ):
        exclude(~valid, f'outside_domain:{column}')
    exclude(pd.to_datetime(df_raw['Tanggal'], errors='coerce').isna(), 'invalid_date')
    exclude(~df_raw['Observasi'].isin(('Y', 'N')), 'invalid_observation_label')
    exclude(~df_raw['Prediksi'].isin(('Y', 'N')), 'missing_model_prediction')
    if 'moon_semidiameter_deg_BT' in df_raw:
        diameter = pd.to_numeric(df_raw['moon_semidiameter_deg_BT'], errors='coerce')
        exclude(~np.isfinite(diameter) | (diameter <= 0), 'invalid_semidiameter')
        df_raw['moon_semidiameter_deg_BT'] = diameter

    exclusions = [{
        'No': row['No'], 'Lokasi': row['Lokasi'], 'reason': reasons.at[index],
    } for index, row in df_raw.loc[reasons.ne('')].iterrows()]
    df_raw = df_raw.loc[reasons.eq('')].copy()
    df = pd.DataFrame(index=df_raw.index)
    for column in ('No', 'Tanggal', 'Lokasi', 'Lat', 'Lon', 'Elv'):
        df[column] = df_raw[column]
    df['Phase Angle (°)'] = df_raw[phase_column]
    df['Lebar Sabit (arcmin)'] = df_raw['W_arcmin']
    df['Sun Alt BT (°)'] = df_raw['sun_alt_BT']
    df['Moon Alt (°)'] = df_raw['Moon_Alt_BT']
    df['Elongasi (°)'] = df_raw['Elongasi_BT']
    df['Sky Bright (nL)'] = df_raw['Sky_Bright_BT']
    df['Lum Hilal (nL)'] = df_raw['Lum_Hilal_BT']
    df['k_V'] = df_raw['kV_BT']
    df['Best Time Tel'] = df_raw['Best_Time_Tel']
    if 'moon_semidiameter_deg_BT' in df_raw:
        df['Moon Semidiameter (deg)'] = df_raw['moon_semidiameter_deg_BT']
    else:
        # Legacy CSV: reconstruct each usable row, recording missing timestamps
        # or invalid geometry without aborting the remaining observations.
        semidiameters = {}
        for index, row in df.iterrows():
            try:
                semidiameters[index] = _observation_semidiameter_deg(row)
            except (ValueError, TypeError, KeyError) as error:
                exclusions.append({'No': row['No'], 'Lokasi': row['Lokasi'],
                                   'reason': f'geometry_reconstruction:{error}'})
        keep = df.index.isin(semidiameters)
        df = df.loc[keep].copy()
        df_raw = df_raw.loc[df.index]
        df['Moon Semidiameter (deg)'] = pd.Series(semidiameters, dtype='float64')
    for key in ('extinction_mag_v', 'transmission_v'):
        csv_key = f'{key}_BT'
        if csv_key in df_raw:
            df[key] = df_raw[csv_key]
    df['RH (%)'] = df_raw['RH_BT']
    df['T (°C)'] = df_raw['T_BT']
    df['Leg_Time_min'] = df_raw['Leg_Time_min']
    df['Δm Tel Opt'] = df_raw['Dm_Tel_BT']
    df['Obs (Y/N)'] = df_raw['Observasi']
    df['Cocok?'] = (df_raw['Prediksi'] == df_raw['Observasi']).apply(
        lambda match: '✓' if match else '✗')
    for column in ('Observation_Method', 'Observation_Source', 'Comparison_Scope'):
        if column in df_raw:
            df[column] = df_raw[column].fillna('unknown').replace('', 'unknown')
        else:
            df[column] = 'unknown'
    for column in ('Actual_Telescope_Config_Available', 'Actual_Observation_Time_Available'):
        df[column] = df_raw[column] if column in df_raw else pd.NA
    df.attrs['excluded_rows'] = exclusions
    df.attrs['comparison_interpretation'] = 'descriptive_sensitivity_not_visual_validation'
    if exclusions:
        counts = pd.Series([item['reason'] for item in exclusions]).value_counts()
        details = ', '.join(f'{reason}={count}' for reason, count in counts.items())
        print(f'  Diagnostic: {len(exclusions)} baris dikeluarkan ({details}).')
    return df.reset_index(drop=True)


# ===================================================================
# FUNGSI UTILITAS
# ===================================================================

def estimate_azimuth_diff(elongation, sun_alt, moon_alt):
    rd = math.pi / 180.0
    cos_elong = math.cos(elongation * rd)
    sin_sun = math.sin(sun_alt * rd)
    sin_moon = math.sin(moon_alt * rd)
    cos_sun = math.cos(sun_alt * rd)
    cos_moon = math.cos(moon_alt * rd)
    denom = cos_sun * cos_moon
    if abs(denom) < 1e-10:
        return elongation
    cos_daz = (cos_elong - sin_sun * sin_moon) / denom
    cos_daz = max(-1.0, min(1.0, cos_daz))
    return math.degrees(math.acos(cos_daz))


def _telescope_factors():
    return TelescopeVisibilityModel().calculate_factors(
        D=TEL_APERTURE, Ds=TEL_OBSTRUCTION, M=TEL_MAG, age=TEL_AGE,
        t1=TEL_TRANS, n=TEL_NSURFACES,
    )


def compute_telescope_Ba(B_sky_nL):
    if not math.isfinite(B_sky_nL) or B_sky_nL < 0:
        raise ValueError("Background harus finite dan non-negatif.")
    return B_sky_nL * _telescope_factors()['surface_brightness_factor']


def _observation_semidiameter_deg(row):
    """Ambil r dari CSV baru, atau turunkan dari DE440s untuk CSV lama."""
    semidiameter = row.get('Moon Semidiameter (deg)')
    if semidiameter is not None and pd.notna(semidiameter):
        semidiameter = float(semidiameter)
        if math.isfinite(semidiameter) and semidiameter > 0:
            return semidiameter
        raise ValueError("Semidiameter pengamatan harus finite dan positif.")
    best_time = row.get('Best Time Tel')
    if best_time is None or pd.isna(best_time) or not str(best_time).strip():
        raise ValueError("CSV memerlukan semidiameter atau Best_Time_Tel untuk geometri DE440s.")
    date = pd.to_datetime(row['Tanggal']).date()
    local_datetime = pd.to_datetime(f'{date} {best_time}').to_pydatetime()
    timezone_str = tentukan_timezone_indonesia(float(row['Lon']))
    local_datetime = timezone(timezone_str).localize(local_datetime)
    location = set_location(float(row['Lat']), float(row['Lon']), float(row['Elv']))
    return moon_semidiameter_time_local(timezone_str, location, local_datetime)


def compute_delta_m(L_nL, B_nL, elongation_deg, mode='naked_eye', *, moon_sd_deg):
    """Margin dengan semidiameter pengamatan eksplisit [derajat]."""
    if mode not in ('naked_eye', 'telescope'):
        raise ValueError("mode harus naked_eye atau telescope.")
    if not math.isfinite(L_nL) or L_nL < 0:
        raise ValueError("Luminansi objek harus finite dan non-negatif.")
    if not math.isfinite(B_nL) or B_nL <= 0:
        raise ValueError("Background harus finite dan positif.")
    B_cd = nL_to_cdm2(B_nL)
    delta_B_obj_cd = nL_to_cdm2(L_nL)

    C_obj = delta_B_obj_cd / B_cd
    A_arcmin2 = crescent_area_arcmin2(elongation_deg, moon_sd_deg)
    A_sr = arcmin2_to_sr(A_arcmin2)
    if A_sr <= 0:
        return {'delta_m': float('-inf'), 'C_obj': C_obj, 'C_th': float('inf')}

    if mode == 'naked_eye':
        C_th = contrast_threshold(A_sr, B_cd, F=F_NAKED, mode='auto')
    else:
        factors = _telescope_factors()
        threshold = telescopic_extended_threshold(
            A_sr, B_cd, TEL_APERTURE / 1000, TEL_MAG,
            p=factors['De'] / 1000, Ft=factors['Ft'], F=F_TEL,
            dimming_factor=factors['surface_brightness_factor'],
        )
        C_th = threshold['C_th']

    delta_m = visibility_margin_mag(C_obj, C_th)

    return {'delta_m': delta_m, 'C_obj': C_obj, 'C_th': C_th}


def compute_full_chain(row, rh_override):
    """Jalankan rantai fisika lengkap dengan RH tertentu.
    Menggunakan sun_alt dari data Skyfield (bukan estimasi).
    """
    moon_alt = row['Moon Alt (°)']
    phase_angle = row['Phase Angle (°)']
    elongation = row['Elongasi (°)']
    temperature = row['T (°C)']
    lat = row['Lat']
    elv = row['Elv']
    moon_sd_deg = _observation_semidiameter_deg(row)

    if not math.isfinite(moon_alt) or not 0 < moon_alt <= 90:
        raise ValueError('Diagnostik memerlukan altitude bulan finite di atas horizon.')
    if not math.isfinite(temperature) or not math.isfinite(elv):
        raise ValueError('Suhu dan elevasi diagnostik harus finite.')
    if not math.isfinite(lat) or not -90 <= lat <= 90:
        raise ValueError('Lintang diagnostik harus finite di -90..90 derajat.')

    tgl = pd.to_datetime(row['Tanggal'])
    month = tgl.month
    year = tgl.year

    # Sun altitude dari data Skyfield (FIX: bukan estimasi lagi)
    sun_alt = row['Sun Alt BT (°)']
    if not math.isfinite(sun_alt) or not -90 <= sun_alt <= 90:
        raise ValueError('Altitude matahari harus finite di -90..90 derajat.')
    azisun = estimate_azimuth_diff(elongation, sun_alt, moon_alt)

    # RH = 100% berada di singularitas rumus aerosol Schaefer.
    # Lewati input ini secara eksplisit; jangan menggantinya dengan atmosfer asumsi.
    if not math.isfinite(rh_override) or not 0 <= rh_override < 100:
        raise ValueError("Diagnostik Schaefer memerlukan RH finite dalam 0 <= RH < 100%.")
    rh_used = rh_override

    # Sky Brightness
    result_sky = hitung_sky_brightness(
        month=month, year=year,
        altsun=sun_alt, azisun=azisun,
        humidity=rh_used, temperature=temperature,
        latitude=lat, elevation=elv,
        alt_objek=moon_alt,
    )
    B_sky_nL = float(result_sky['sky_brightness'])
    k_v = float(result_sky['k_v'])
    extinction_mag_v = float(result_sky['extinction_mag_v'])
    transmission_v = float(result_sky['transmission_v'])

    if B_sky_nL <= 0:
        raise ValueError("Sky brightness Schaefer harus positif.")

    # Luminansi Hilal
    L_star_s10 = hitung_luminansi_intrinsik(phase_angle, elongation, moon_sd_deg)
    L_hilal_nL = terapkan_transmisi_atmosfer(L_star_s10, transmission_v)

    # Visibilitas
    result_ne = compute_delta_m(L_hilal_nL, B_sky_nL, elongation, 'naked_eye', moon_sd_deg=moon_sd_deg)
    result_tel = compute_delta_m(L_hilal_nL, B_sky_nL, elongation, 'telescope', moon_sd_deg=moon_sd_deg)

    return {
        'rh_used': rh_used, 'k_v': k_v,
        'extinction_mag_v': extinction_mag_v, 'transmission_v': transmission_v,
        'L_star_s10': L_star_s10,
        'moon_semidiameter': moon_sd_deg,
        'B_sky_nL': B_sky_nL, 'L_hilal_nL': L_hilal_nL,
        'C_obj': result_ne['C_obj'],
        'C_th_ne': result_ne['C_th'], 'C_th_tel': result_tel['C_th'],
        'dm_ne': result_ne['delta_m'], 'dm_tel': result_tel['delta_m'],
        'sun_alt': sun_alt,
    }


def classify_obs_type(obs_yn, cocok):
    if obs_yn == 'Y' and cocok == '✓': return 'TP'
    elif obs_yn == 'Y': return 'FN'
    elif obs_yn == 'N' and cocok == '✓': return 'TN'
    else: return 'FP'


def categorize_fn(butuh_delta_rh):
    if not isinstance(butuh_delta_rh, (int, float)): return 'C'
    abs_drh = abs(butuh_delta_rh)
    if abs_drh <= 5: return 'A'
    elif abs_drh <= 20: return 'B'
    else: return 'C'


# ===================================================================
# 4.5.1 — SENSITIVITAS OAT
# ===================================================================

def analyze_rh_sensitivity(df):
    print("\n" + "=" * 70)
    print("4.5.1  SENSITIVITAS OAT Dm TERHADAP VARIASI RH")
    print("Kecocokan label/FN/FP bersifat deskriptif lintas metode atau metode tidak diketahui.")
    print("=" * 70)

    all_rows = []
    critical_rows = []

    for idx, row in df.iterrows():
        obs_no = int(row['No'])
        rh_baseline = row['RH (%)']
        obs_yn = row['Obs (Y/N)']
        tipe = classify_obs_type(obs_yn, row['Cocok?'])

        print(f"\n  Obs #{obs_no:2d} {row['Lokasi'][:30]:30s} ({tipe}) RH={rh_baseline:.1f}%")

        dm_at_deltas = {}
        kv_at_deltas = {}

        for delta_rh in RH_DELTAS:
            rh_test = rh_baseline + delta_rh
            if rh_test < 0 or rh_test >= 100: continue

            result = compute_full_chain(row, rh_test)
            dm_at_deltas[delta_rh] = result['dm_tel']
            kv_at_deltas[delta_rh] = result['k_v']

            all_rows.append({
                **_observation_provenance(row),
                'No': obs_no, 'Lokasi': row['Lokasi'],
                'Obs': obs_yn, 'Tipe': tipe,
                'RH_baseline': rh_baseline, 'ΔRH': delta_rh,
                'RH_test': rh_test, 'k_V': result['k_v'],
                'extinction_mag_v': result['extinction_mag_v'],
                'transmission_v': result['transmission_v'],
                'moon_semidiameter_deg': result['moon_semidiameter'],
                'B_sky_nL': result['B_sky_nL'], 'L_hilal_nL': result['L_hilal_nL'],
                'Δm_NE': result['dm_ne'], 'Δm_Tel': result['dm_tel'],
                'C_obj': result['C_obj'],
            })

        # RH kritis (interpolasi)
        deltas_sorted = sorted(dm_at_deltas.keys())
        critical_rh = None
        for i in range(len(deltas_sorted) - 1):
            d1, d2 = deltas_sorted[i], deltas_sorted[i+1]
            dm1, dm2 = dm_at_deltas[d1], dm_at_deltas[d2]
            if dm1 <= -90 or dm2 <= -90: continue
            if (dm1 <= 0 <= dm2) or (dm2 <= 0 <= dm1):
                frac = -dm1 / (dm2 - dm1) if (dm2 - dm1) != 0 else 0
                critical_rh = rh_baseline + d1 + frac * (d2 - d1)
                break

        dm_base = dm_at_deltas.get(0, -99)
        dm_minus20 = dm_at_deltas.get(-20, dm_at_deltas.get(-25, -99))
        improvement = dm_minus20 - dm_base if dm_base > -90 and dm_minus20 > -90 else 0
        butuh_drh = (critical_rh - rh_baseline) if critical_rh else 'N/A'
        kategori = categorize_fn(butuh_drh) if tipe == 'FN' else '-'

        critical_rows.append({
            **_observation_provenance(row),
            'No': obs_no, 'Lokasi': row['Lokasi'],
            'Obs': obs_yn, 'Tipe': tipe,
            'RH_baseline': rh_baseline, 'k_V_baseline': kv_at_deltas.get(0, 0),
            'Δm_Tel_baseline': dm_base, 'Δm_Tel_RH-20': dm_minus20,
            'Improvement_20': improvement,
            'RH_kritis': critical_rh if critical_rh else 'N/A',
            'Butuh_ΔRH': butuh_drh, 'Kategori_FN': kategori,
        })

        status = f"RH_kritis={critical_rh:.1f}%" if critical_rh else "di luar range"
        kat_str = f"  [{kategori}]" if tipe == 'FN' else ""
        print(f"    Δm(base)={dm_base:+.2f}  Δm(RH-20)={dm_minus20:+.2f}  {status}{kat_str}")

    critical_df = pd.DataFrame(critical_rows)

    fn_rows = critical_df[critical_df['Tipe'] == 'FN']
    if len(fn_rows) > 0:
        print("\n  KELOMPOK SENSITIVITAS FN LINTAS METODE (BUKAN DIAGNOSIS TERVALIDASI):")
        for cat_label, desc in [('A','perubahan RH ≤5pp'), ('B','perubahan RH 5-20pp'),
                                ('C','perubahan RH >20pp atau ambang tidak tercapai')]:
            n = sum(fn_rows['Kategori_FN'] == cat_label)
            print(f"    {cat_label} ({desc}): {n}/{len(fn_rows)}")

    return pd.DataFrame(all_rows), critical_df


# ===================================================================
# 4.5.2 — DEKOMPOSISI JALUR ERROR
# ===================================================================

def analyze_error_decomposition(df):
    print("\n" + "=" * 70)
    print("4.5.2  DEKOMPOSISI SENSITIVITAS (KASTNER vs SCHAEFER)")
    print("=" * 70)

    fn_fp = df[((df['Obs (Y/N)']=='Y')&(df['Cocok?']!='✓'))|
               ((df['Obs (Y/N)']=='N')&(df['Cocok?']!='✓'))]
    if len(fn_fp) == 0:
        print("  Tidak ada FN/FP."); return pd.DataFrame()

    decomp_rows = []
    for _, row in fn_fp.iterrows():
        obs_no = int(row['No'])
        rh_base = row['RH (%)']
        tipe = classify_obs_type(row['Obs (Y/N)'], row['Cocok?'])

        res_base = compute_full_chain(row, rh_base)
        rh_low = max(rh_base - 20, 0.0)
        res_low = compute_full_chain(row, rh_low)

        L_low_transmission = terapkan_transmisi_atmosfer(res_base['L_star_s10'], res_low['transmission_v'])
        res_only_L = compute_delta_m(L_low_transmission, res_base['B_sky_nL'], row['Elongasi (°)'], 'telescope',
                                    moon_sd_deg=res_base['moon_semidiameter'])
        res_only_B = compute_delta_m(res_base['L_hilal_nL'], res_low['B_sky_nL'], row['Elongasi (°)'], 'telescope',
                                    moon_sd_deg=res_base['moon_semidiameter'])

        dm_full = res_low['dm_tel'] - res_base['dm_tel']
        dm_L = res_only_L['delta_m'] - res_base['dm_tel']
        dm_B = res_only_B['delta_m'] - res_base['dm_tel']
        dm_int = dm_full - dm_L - dm_B

        decomp_rows.append({
            **_observation_provenance(row),
            'No': obs_no, 'Lokasi': row['Lokasi'], 'Tipe': tipe,
            'RH_base': rh_base, 'RH_low': rh_low,
            'k_V_base': res_base['k_v'], 'k_V_low': res_low['k_v'],
            'extinction_mag_v_base': res_base['extinction_mag_v'],
            'extinction_mag_v_low': res_low['extinction_mag_v'],
            'transmission_v_base': res_base['transmission_v'],
            'transmission_v_low': res_low['transmission_v'],
            'Δm_base': res_base['dm_tel'], 'Δm_full_change': dm_full,
            'Δm_dari_L (Kastner)': dm_L, 'Δm_dari_B (Schaefer)': dm_B,
            'Δm_interaksi': dm_int,
            'Kontribusi_L (%)': abs(dm_L)/abs(dm_full)*100 if abs(dm_full)>0.001 else 0,
            'Kontribusi_B (%)': abs(dm_B)/abs(dm_full)*100 if abs(dm_full)>0.001 else 0,
        })

        print(f"\n  #{obs_no} ({tipe}): Kastner={dm_L:+.3f} Schaefer={dm_B:+.3f} Int={dm_int:+.3f}")

    decomp_df = pd.DataFrame(decomp_rows)
    if len(decomp_df) > 0:
        print(f"\n  Rata-rata: Kastner={decomp_df['Kontribusi_L (%)'].mean():.1f}%  "
              f"Schaefer={decomp_df['Kontribusi_B (%)'].mean():.1f}%")
    return decomp_df


# ===================================================================
# 4.5.3 — ERROR BAR
# ===================================================================

def analyze_error_bars(sensitivity_df, critical_df):
    print("\n" + "=" * 70)
    print("4.5.3  ERROR BAR Dm DARI KETIDAKPASTIAN ERA5")
    print("=" * 70)

    RH_UNC = [5, 10, 15]
    rows = []
    for _, cr in critical_df.iterrows():
        obs_no = cr['No']
        dm_base = cr['Δm_Tel_baseline']
        tipe = cr['Tipe']
        obs_sens = sensitivity_df[sensitivity_df['No'] == obs_no]

        for unc in RH_UNC:
            dm_low = obs_sens[obs_sens['ΔRH'] == -unc]
            dm_high = obs_sens[obs_sens['ΔRH'] == unc]
            dm_low_val = dm_low['Δm_Tel'].values[0] if len(dm_low) > 0 else dm_base
            dm_high_val = dm_high['Δm_Tel'].values[0] if len(dm_high) > 0 else dm_base
            error_bar = (dm_low_val - dm_high_val) / 2.0
            covers = (min(dm_low_val, dm_high_val) <= 0 <= max(dm_low_val, dm_high_val))
            rows.append({
                **_observation_provenance(cr),
                'No': obs_no, 'Lokasi': cr['Lokasi'], 'Obs': cr['Obs'], 'Tipe': tipe,
                'Δm_baseline': dm_base, 'σ_RH (±pp)': unc,
                'Δm_low_RH': dm_low_val, 'Δm_high_RH': dm_high_val,
                'Error_bar (±mag)': abs(error_bar),
                'Mencakup_Δm=0': 'Ya' if covers else 'Tidak',
            })

    result = pd.DataFrame(rows)
    print("\n  Sensitivitas FN lintas metode: ambang model tercakup oleh variasi RH:")
    for unc in RH_UNC:
        fn_sub = result[(result['σ_RH (±pp)']==unc) & (result['Tipe']=='FN')]
        if len(fn_sub) > 0:
            n_cov = sum(fn_sub['Mencakup_Δm=0']=='Ya')
            print(f"    ±{unc:2d} pp: {n_cov}/{len(fn_sub)} FN ({n_cov/len(fn_sub)*100:.0f}%)")
    return result


# ===================================================================
# 4.5.4 — ERA5 vs MERRA-2
# ===================================================================

def analyze_era5_vs_merra2(df_era5, df_merra2):
    print("\n" + "=" * 70)
    print("4.5.4  PERBANDINGAN ERA5 vs MERRA-2")
    print("=" * 70)

    provenance_era5 = [column for column in PROVENANCE_COLUMNS if column in df_era5]
    provenance_merra2 = [column for column in PROVENANCE_COLUMNS if column in df_merra2]
    merged = pd.merge(
        df_era5[['No','Lokasi','Obs (Y/N)','Cocok?','RH (%)','k_V','T (°C)','Δm Tel Opt',
                 *provenance_era5]],
        df_merra2[['No','RH (%)','k_V','T (°C)','Δm Tel Opt','Cocok?', *provenance_merra2]],
        on='No', suffixes=('_ERA5','_MERRA2'))

    merged['ΔRH'] = merged['RH (%)_MERRA2'] - merged['RH (%)_ERA5']
    merged['Δk_V'] = merged['k_V_MERRA2'] - merged['k_V_ERA5']
    merged['ΔΔm'] = merged['Δm Tel Opt_MERRA2'] - merged['Δm Tel Opt_ERA5']

    print(f"\n  Observasi: {len(merged)}")
    for var, label in [('ΔRH','RH(pp)'), ('Δk_V','k_V'), ('ΔΔm','Δm Tel')]:
        v = merged[var].dropna()
        print(f"    {label:8s}: mean={v.mean():+.3f} std={v.std():.3f} [{v.min():+.3f}, {v.max():+.3f}]")

    from scipy import stats as sp
    dm_e = merged['Δm Tel Opt_ERA5'].values
    dm_m = merged['Δm Tel Opt_MERRA2'].values
    valid = (dm_e > -90) & (dm_m > -90)
    if sum(valid) >= 3:
        r_p, p_p = sp.pearsonr(dm_e[valid], dm_m[valid])
        r_s, p_s = sp.spearmanr(dm_e[valid], dm_m[valid])
        print(f"\n  Pearson r={r_p:.4f} (p={p_p:.2e})")
        print(f"  Spearman ρ={r_s:.4f} (p={p_s:.2e})")

    # CM
    obs = merged['Obs (Y/N)'].values
    for label, cocok_col in [('ERA5','Cocok?_ERA5'), ('MERRA2','Cocok?_MERRA2')]:
        cocok = merged[cocok_col].values
        tp = sum((obs=='Y')&(cocok=='✓')); fn = sum((obs=='Y')&(cocok!='✓'))
        tn = sum((obs=='N')&(cocok=='✓')); fp = sum((obs=='N')&(cocok!='✓'))
        n = tp+fn+tn+fp
        agreement = f'{(tp+tn)/n:.1%}' if n else 'N/A'
        print(f"\n  {label}: TP={tp} FN={fn} FP={fp} TN={tn}; "
              f"kecocokan lintas metode={agreement} (deskriptif)")

    return merged


# ===================================================================
# OUTPUT: EXCEL + PLOTS
# ===================================================================

def save_results(sens_df, crit_df, decomp_df, eb_df, m2_df, filepath, input_audit=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    sty = {
        'hdr': Font(name='Arial', bold=True, size=10, color='FFFFFF'),
        'hdr_fill': PatternFill('solid', fgColor='1F4E79'),
        'data': Font(name='Arial', size=10),
        'center': Alignment(horizontal='center', vertical='center'),
        'border': Border(left=Side('thin','B0B0B0'), right=Side('thin','B0B0B0'),
                         top=Side('thin','B0B0B0'), bottom=Side('thin','B0B0B0')),
        'fn_fill': PatternFill('solid', fgColor='FFC7CE'),
    }

    def write_sheet(ws, df, name):
        ws.title = name
        for ci, h in enumerate(df.columns, 1):
            c = ws.cell(row=1, column=ci, value=h)
            c.font = sty['hdr']; c.fill = sty['hdr_fill']
            c.alignment = sty['center']; c.border = sty['border']
        for ri, (_, row) in enumerate(df.iterrows(), 2):
            for ci, col in enumerate(df.columns, 1):
                val = row[col]
                if pd.isna(val):
                    val = None
                elif isinstance(val, float):
                    if abs(val) > 1e6: val = f"{val:.4e}"
                    elif abs(val) < 0.001 and val != 0: val = f"{val:.6f}"
                    else: val = round(val, 4)
                c = ws.cell(row=ri, column=ci, value=val)
                c.font = sty['data']; c.alignment = sty['center']; c.border = sty['border']
        for ci in range(1, len(df.columns)+1):
            ml = max(len(str(df.columns[ci-1])),
                     max((len(str(ws.cell(row=r,column=ci).value or ''))
                          for r in range(2, min(len(df)+2,50))), default=8))
            ws.column_dimensions[get_column_letter(ci)].width = min(ml+2, 25)
        ws.freeze_panes = 'A2'

    # Pivot sensitivitas
    pivot = []
    for obs_no in sorted(sens_df['No'].unique()):
        od = sens_df[sens_df['No']==obs_no]
        base = od[od['ΔRH']==0]
        if len(base)==0: continue
        b = base.iloc[0]
        p = {**_observation_provenance(b),
             'No':obs_no, 'Lokasi':b['Lokasi'], 'Obs':b['Obs'], 'Tipe':b['Tipe'],
             'RH_base':b['RH_baseline'], 'k_V_base':b['k_V'],
             'extinction_mag_v_base':b['extinction_mag_v'],
             'transmission_v_base':b['transmission_v']}
        for d in RH_DELTAS:
            dd = od[od['ΔRH']==d]
            if len(dd)>0:
                p[f'Δm_Tel({d:+d})'] = round(dd.iloc[0]['Δm_Tel'], 3)
        pivot.append(p)

    ws1 = wb.active; write_sheet(ws1, pd.DataFrame(pivot), "4.5.1 Sensitivitas RH")
    ws2 = wb.create_sheet(); write_sheet(ws2, crit_df, "4.5.1 RH Kritis")
    ws3 = wb.create_sheet()
    if len(decomp_df)>0: write_sheet(ws3, decomp_df, "4.5.2 Dekomposisi")
    else: ws3.title = "4.5.2 Dekomposisi"; ws3.cell(1,1,value="Tidak ada FN/FP")
    ws4 = wb.create_sheet(); write_sheet(ws4, eb_df, "4.5.3 Error Bar")
    ws5 = wb.create_sheet()
    if m2_df is not None and len(m2_df)>0: write_sheet(ws5, m2_df, "4.5.4 ERA5 vs MERRA2")
    else: ws5.title = "4.5.4 ERA5 vs MERRA2"
    ws6 = wb.create_sheet(); write_sheet(ws6, sens_df, "Data Detail")
    ws7 = wb.create_sheet('Cakupan Perbandingan')
    for label, value in (
        ('Interpretasi', 'Sensitivitas dan kecocokan label lintas metode atau metode tidak diketahui'),
        ('Validasi visual empiris', 'Tidak dilakukan oleh analisis diagnostik ini'),
        ('F naked eye referensi', F_NAKED), ('F teleskop referensi', F_TEL),
        ('Kategori FN', 'Kelompok perubahan RH; bukan diagnosis bias/struktur yang tervalidasi'),
        ('Metode/sumber pengamatan', 'Dipertahankan dari input; metadata yang hilang tetap unknown'),
    ):
        ws7.append((label, value))
    ws7.column_dimensions['A'].width = 30
    ws7.column_dimensions['B'].width = 95
    audit_rows = []
    for source, audit in (input_audit or {}).items():
        audit_rows.extend({'Input': source, **item} for item in audit.get('excluded_rows', []))
    if audit_rows:
        write_sheet(wb.create_sheet(), pd.DataFrame(audit_rows), 'Input Dikeluarkan')

    os.makedirs(os.path.dirname(filepath) or '.', exist_ok=True)
    wb.save(filepath)
    print(f"\n  Saved: {filepath}")


def plot_sensitivity(sens_df, crit_df, filepath):
    try:
        import matplotlib; matplotlib.use('Agg')
        import matplotlib.pyplot as plt
    except ImportError: return

    misclass = sens_df[sens_df['Tipe'].isin(['FN','FP'])]
    if len(misclass)==0: return

    fig, axes = plt.subplots(1, 2, figsize=(16, 7))
    ax1 = axes[0]
    obs_list = sorted(misclass['No'].unique())
    colors = plt.cm.tab10(range(len(obs_list)))

    for i, obs_no in enumerate(obs_list):
        od = misclass[misclass['No']==obs_no].sort_values('ΔRH')
        t = od.iloc[0]['Tipe']
        ax1.plot(od['ΔRH'], od['Δm_Tel'], marker='*' if t=='FN' else 's',
                 markersize=6, color=colors[i], linewidth=2,
                 label=f"#{obs_no} {od.iloc[0]['Lokasi'][:20]} ({t})")
    ax1.axhline(y=0, color='green', linewidth=2, linestyle='--', alpha=0.8)
    ax1.set_xlabel('Perubahan RH (pp)'); ax1.set_ylabel('Δm Tel (mag)')
    ax1.set_title('4.5.1  Sensitivitas Δm vs ΔRH'); ax1.legend(fontsize=7); ax1.grid(True, alpha=0.3)

    ax2 = axes[1]
    all_base = sens_df[sens_df['ΔRH']==0]
    for t, c, m, s in [('TN','blue','o',50),('TP','green','D',70),('FN','red','*',100),('FP','orange','s',70)]:
        sub = all_base[all_base['Tipe']==t]
        if len(sub)>0:
            ax2.scatter(sub['k_V'], sub['Δm_Tel'], c=c, s=s, alpha=0.8, marker=m,
                        label=f'{t} (n={len(sub)})', edgecolors='black', linewidths=0.5)
    ax2.axhline(y=0, color='green', linewidth=2, linestyle='--', alpha=0.8)
    ax2.set_xlabel('k_V'); ax2.set_ylabel('Δm Tel (mag)')
    ax2.set_title('k_V vs Δm; FN/FP lintas metode (deskriptif)'); ax2.legend(); ax2.grid(True, alpha=0.3)

    plt.tight_layout(); fig.savefig(filepath, dpi=150, bbox_inches='tight'); plt.close(fig)
    print(f"  Saved: {filepath}")


def plot_decomposition(decomp_df, filepath):
    try:
        import matplotlib; matplotlib.use('Agg')
        import matplotlib.pyplot as plt
    except ImportError: return
    if len(decomp_df)==0: return

    fig, ax = plt.subplots(figsize=(12, 6))
    labels = [f"#{int(r['No'])} {r['Lokasi'][:18]}" for _, r in decomp_df.iterrows()]
    x = range(len(labels)); w = 0.25

    ax.bar([i-w for i in x], decomp_df['Δm_dari_L (Kastner)'], w, label='Kastner', color='#E74C3C')
    ax.bar(x, decomp_df['Δm_dari_B (Schaefer)'], w, label='Schaefer', color='#3498DB')
    ax.bar([i+w for i in x], decomp_df['Δm_interaksi'], w, label='Interaksi', color='#95A5A6')

    ax.set_xticks(x); ax.set_xticklabels(labels, rotation=30, ha='right', fontsize=9)
    ax.set_xlabel('Observasi'); ax.set_ylabel('Kontribusi (mag)')
    ax.set_title('4.5.2  Dekomposisi: Kastner vs Schaefer')
    ax.legend(); ax.grid(True, axis='y', alpha=0.3); ax.axhline(y=0, color='black', linewidth=0.5)

    plt.tight_layout(); fig.savefig(filepath, dpi=150, bbox_inches='tight'); plt.close(fig)
    print(f"  Saved: {filepath}")


# ===================================================================
# MAIN
# ===================================================================

def main():
    configure_console_encoding()
    print("\n" + "#" * 70)
    print("  ANALISIS DIAGNOSTIK CRUMEY — Tahap 4.5")
    print("  Sensitivitas/deskriptif lintas metode; tidak mengkalibrasi threshold visual.")
    print(f"  F_NAKED={F_NAKED}, F_TEL={F_TEL}, TEL_AGE={TEL_AGE}")
    print("#" * 70)

    if not os.path.exists(INPUT_ERA5):
        print(f"\n  [!] File tidak ditemukan: {INPUT_ERA5}"); return

    print(f"\n  Membaca ERA5: {INPUT_ERA5}")
    df_era5 = load_observation_data(INPUT_ERA5)
    print(f"  {len(df_era5)} observasi")
    if df_era5.empty:
        print('  Tidak ada input fisik yang valid untuk analisis sensitivitas.')
        return

    df_merra2 = None
    if os.path.exists(INPUT_MERRA2):
        print(f"  Membaca MERRA-2: {INPUT_MERRA2}")
        df_merra2 = load_observation_data(INPUT_MERRA2)
        print(f"  {len(df_merra2)} observasi")

    sens_df, crit_df = analyze_rh_sensitivity(df_era5)
    decomp_df = analyze_error_decomposition(df_era5)
    eb_df = analyze_error_bars(sens_df, crit_df)

    m2_df = None
    if df_merra2 is not None and not df_merra2.empty:
        m2_df = analyze_era5_vs_merra2(df_era5, df_merra2)

    # Simpan
    excel_path = os.path.join(OUTPUT_DIR, "Analisis_Diagnostik_Crumey.xlsx")
    input_audit = {'ERA5': df_era5.attrs}
    if df_merra2 is not None:
        input_audit['MERRA2'] = df_merra2.attrs
    save_results(sens_df, crit_df, decomp_df, eb_df, m2_df, excel_path, input_audit=input_audit)

    plot_sensitivity(sens_df, crit_df, os.path.join(OUTPUT_DIR, "Sensitivitas_RH_dan_kV.png"))
    plot_decomposition(decomp_df, os.path.join(OUTPUT_DIR, "Dekomposisi_Jalur_Error.png"))

    print(f"\n{'#' * 70}")
    print("  ANALISIS 4.5 SELESAI")
    print(f"{'#' * 70}\n")


if __name__ == "__main__":
    main()
