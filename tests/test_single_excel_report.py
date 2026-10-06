"""Small end-to-end checks for the single-location Excel modelling chain."""

import contextlib
from copy import deepcopy
import io
import math
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'Core'))
from core_crescent_visibility import HilalVisibilityCalculator
from full_rumus_crumey import mag_to_lux


class SingleExcelReportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.calc = HilalVisibilityCalculator(
            'Excel Semarang', -6.917, 110.348, 89, 'Asia/Jakarta', 9, 1444,
            sumber_atmosfer='manual', manual_rh=75, manual_t=25,
        )
        with contextlib.redirect_stdout(io.StringIO()):
            cls.calc.jalankan_perhitungan_lengkap(
                mode='optimal', interval_menit=5, aperture=120, magnification=40,
                transmission=0.9, n_surfaces=4, central_obstruction=30,
                pupil_diameter_mm=4, field_factor=2.2,
            )

    def export(self, calc=None):
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        path = Path(folder.name) / 'nested' / 'single.xlsx'
        calc = calc or self.calc
        # Export must use captured results, without weather/model recomputation.
        with (patch.object(calc, '_fetch_atmosfer', side_effect=AssertionError('weather called')),
              patch.object(calc, 'hitung_visibilitas_naked_eye', side_effect=AssertionError('model called')),
              contextlib.redirect_stdout(io.StringIO())):
            calc.simpan_ke_excel(str(path))
        self.assertTrue(Path(str(path) + '.atmosphere.json').is_file())
        wb = load_workbook(path, data_only=True)
        self.addCleanup(wb.close)
        return wb

    @staticmethod
    def row(ws, symbol):
        return next(row for row in ws.iter_rows() if row[1].value == symbol)

    def test_layout_keeps_methods_and_optima_in_distinct_numeric_columns(self):
        wb = self.export()
        for name in ('Ringkasan', 'Input & Konfigurasi', 'Rantai Model', 'Atmosfer', 'Timestep Data', 'Info Program'):
            self.assertIn(name, wb.sheetnames)
        ws = wb['Rantai Model']
        self.assertEqual([ws.cell(5, c).value for c in range(4, 8)],
                         ['Mata telanjang', 'Teleskop', 'Mata telanjang', 'Teleskop'])
        self.assertEqual(ws.freeze_panes, 'D6')
        for col, scene, key in (
            (4, self.calc.hasil, 'delta_m_ne'), (5, self.calc.hasil, 'delta_m_tel'),
            (6, self.calc.hasil['optimal_result_ne'], 'delta_m_ne'),
            (7, self.calc.hasil['optimal_result_tel'], 'delta_m_tel'),
        ):
            cell = self.row(ws, 'Δm')[col - 1]
            self.assertEqual(cell.data_type, 'n')
            self.assertAlmostEqual(cell.value, scene[key], places=12)
        times = self.row(ws, 't')
        self.assertEqual(times[5].value, self.calc.hasil['optimal_time_ne'].replace(tzinfo=None))
        self.assertEqual(times[6].value, self.calc.hasil['optimal_time_tel'].replace(tzinfo=None))

    def test_exported_photometry_and_optics_satisfy_physical_identities(self):
        ws = self.export()['Rantai Model']
        val = lambda symbol, col: self.row(ws, symbol)[col - 1].value
        for col in range(4, 8):
            mv, av = val('M_v', col), val('A_V = DM[2]', col)
            flux = val('I_obs = L × A', col)
            self.assertAlmostEqual(flux / (mag_to_lux(mv) * 10**(-0.4 * av)), 1, places=12)
            self.assertAlmostEqual(val('T_V', col), 10**(-0.4 * av), places=14)
            self.assertTrue(math.isclose(val('L_app', col), val('g', col) * val('L', col), rel_tol=1e-14))
            self.assertTrue(math.isclose(val('B_app', col), val('g', col) * val('B_sky', col), rel_tol=1e-14))
            self.assertAlmostEqual(val('Δm', col), 2.5 * math.log10(val('C_obj/C_th', col)), places=12)
        area_sr = next(row[3].value for row in ws.iter_rows()
                       if row[1].value == 'A' and row[2].value == 'sr')
        self.assertAlmostEqual(val('A_app = M² A', 5) / area_sr, 40**2, places=10)
        self.assertAlmostEqual(val('φ = F F_T F_M', 5), 2.2 * math.sqrt(2), places=12)

    def test_atmospheric_breakdown_matches_total_extinction_and_sky(self):
        ws = self.export()['Atmosfer']
        rows = {row[1].value: row for row in ws.iter_rows() if row[1].value}
        # k_* symbols repeat by band; take the original captured V coefficients
        # and check the independently exported aggregate + sky decomposition.
        scenes = [self.calc.hasil, self.calc.hasil['optimal_result_ne'], self.calc.hasil['optimal_result_tel']]
        for col, scene in zip((4, 5, 6), scenes):
            v = scene['model_trace']['atmosphere']['diagnostics']['bands']['V']
            x = scene['model_trace']['atmosphere']['diagnostics']
            expected = v['k_R'] * x['X_G'] + v['k_A'] * x['X_A'] + v['k_O'] * x['X_O'] + v['k_W'] * x['X_G']
            self.assertAlmostEqual(rows['DM[V]'][col - 1].value, expected, places=12)
            total = rows['B_night_nL'][col - 1].value + min(rows['B_twilight_nL'][col - 1].value, rows['B_daylight_nL'][col - 1].value)
            self.assertAlmostEqual(total / rows['B_sky'][col - 1].value, 1, places=12)

    def test_disabled_telescope_and_unavailable_optima_are_blank(self):
        calc = deepcopy(self.calc)
        calc.hasil.update(use_telescope=False, optimal_result_ne=None, optimal_result_tel=None)
        wb = self.export(calc)
        for name in ('Ringkasan', 'Rantai Model'):
            row = self.row(wb[name], 'Δm')
            self.assertTrue(all(row[c - 1].value is None for c in (5, 6, 7)))
        ws = wb['Timestep Data']
        col = [cell.value for cell in ws[1]].index('Margin Teleskop') + 1
        self.assertIsNone(ws.cell(2, col).value)
        self.assertEqual(len(wb['Ringkasan']._charts[0].series), 2)

    def test_invalid_scenes_have_geometry_but_no_synthetic_margin(self):
        calc = deepcopy(self.calc)
        calc.hasil['all_timestep_results'] = [{
            'waktu_local': calc.hasil['sunset_local'], 'moon_alt': -1, 'sun_alt': -3,
            'valid': False, 'delta_m_ne': -99, 'delta_m_tel': -99, 'rh': 75,
        }]
        ws = self.export(calc)['Timestep Data']
        headers = [cell.value for cell in ws[1]]
        self.assertEqual(ws.cell(2, headers.index('Moon Alt (°)') + 1).value, -1)
        for label in ('Δm Naked Eye', 'Margin Teleskop', 'NE | Keputusan', 'TEL | Keputusan'):
            self.assertIsNone(ws.cell(2, headers.index(label) + 1).value)

    def test_sunset_mode_and_literal_user_text(self):
        calc = deepcopy(self.calc)
        calc.nama_tempat = '=literal location'
        calc.hasil.update(mode='sunset', all_timestep_results=[])
        wb = self.export(calc)
        self.assertFalse(wb['Ringkasan']._charts)
        self.assertIsNone(self.row(wb['Rantai Model'], 'Δm')[5].value)
        self.assertEqual(self.row(wb['Input & Konfigurasi'], 'nama')[3].data_type, 's')

    def test_nonfinite_margins_preserved_and_chart_omits_them(self):
        calc = deepcopy(self.calc)
        for scene in [calc.hasil, *calc.hasil['all_timestep_results'],
                      calc.hasil['optimal_result_ne'], calc.hasil['optimal_result_tel']]:
            if scene:
                scene['delta_m_ne'] = scene['delta_m_tel'] = -math.inf
        wb = self.export(calc)
        self.assertEqual(self.row(wb['Rantai Model'], 'Δm')[3].value, '-inf')
        data = wb['_Data Grafik']
        self.assertTrue(all(row[1].value is None and row[2].value is None for row in data.iter_rows(min_row=2)))


if __name__ == '__main__':
    unittest.main()
