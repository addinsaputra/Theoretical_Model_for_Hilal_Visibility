"""Diagnostic exclusions and observation-method provenance, without APIs.

These tests check safe analysis contracts, not empirical crescent calibration.
"""

import contextlib
import io
import math
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

import pandas as pd
from openpyxl import load_workbook

import hilal_visibility.studies.diagnostics as diagnostic
class DiagnosticProvenanceTests(unittest.TestCase):
    @staticmethod
    def raw_row(no=1):
        return dict(
            No=no, Tanggal='2023-03-22', Lokasi='Diagnostic fixture',
            Lat=-7, Lon=110, Elv=89, Phase_Angle=170, Phase_Angle_BT=170,
            W_arcmin=0.1, sun_alt_BT=-1, Moon_Alt_BT=5, Elongasi_BT=8,
            Sky_Bright_BT=1000, Lum_Hilal_BT=10, kV_BT=0.2,
            Best_Time_Tel='18:00:00', moon_semidiameter_deg_BT=0.26,
            RH_BT=75, T_BT=25, Leg_Time_min=30, Dm_Tel_BT=-1,
            Observasi='Y', Prediksi='N', Status='valid',
            Observation_Method='ccd', Observation_Source='fixture-gallery',
            Comparison_Scope='cross_method_descriptive',
            Actual_Telescope_Config_Available=False,
            Actual_Observation_Time_Available=False,
        )

    def load_rows(self, rows):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'observations.csv'
            pd.DataFrame(rows).to_csv(path, index=False)
            with contextlib.redirect_stdout(io.StringIO()) as output:
                result = diagnostic.load_observation_data(str(path))
        return result, output.getvalue()

    def test_failed_weather_and_nonfinite_inputs_are_excluded_and_audited(self):
        valid = self.raw_row(2)
        failed = {**self.raw_row(1), 'Status': 'invalid', 'Error': 'weather offline',
                  'moon_semidiameter_deg_BT': math.nan, 'Best_Time_Tel': None}
        bad_rh = {**self.raw_row(3), 'RH_BT': math.nan}
        no_source = {**self.raw_row(4), 'Lum_Hilal_BT': 0, 'Dm_Tel_BT': -math.inf}
        with patch.object(diagnostic, '_observation_semidiameter_deg') as reconstruction:
            loaded, output = self.load_rows([failed, valid, bad_rh, no_source])
        reconstruction.assert_not_called()
        self.assertEqual(loaded['No'].tolist(), [2, 4])
        excluded = loaded.attrs['excluded_rows']
        self.assertEqual([item['No'] for item in excluded], [1, 3])
        self.assertEqual([item['reason'] for item in excluded],
                         ['status_invalid', 'nonfinite_input:RH_BT'])
        self.assertIn('2 baris dikeluarkan', output)
        self.assertEqual(loaded.iloc[1]['Δm Tel Opt'], -math.inf)
        for column in diagnostic.PROVENANCE_COLUMNS:
            self.assertEqual(loaded.iloc[0][column], valid[column])

    def test_missing_method_metadata_stays_unknown(self):
        row = self.raw_row()
        for column in diagnostic.PROVENANCE_COLUMNS:
            row.pop(column)
        loaded, _ = self.load_rows([row])
        for column in diagnostic.PROVENANCE_COLUMNS[:3]:
            self.assertEqual(loaded.iloc[0][column], 'unknown')
        for column in diagnostic.PROVENANCE_COLUMNS[3:]:
            self.assertTrue(pd.isna(loaded.iloc[0][column]))

    def test_legacy_geometry_reconstructs_only_usable_rows(self):
        valid = self.raw_row(2)
        failed = {**self.raw_row(1), 'Status': 'invalid', 'Best_Time_Tel': None}
        valid.pop('moon_semidiameter_deg_BT')
        failed.pop('moon_semidiameter_deg_BT')
        with patch.object(diagnostic, '_observation_semidiameter_deg', return_value=0.27) as reconstruct:
            loaded, _ = self.load_rows([failed, valid])
        reconstruct.assert_called_once()
        self.assertEqual(reconstruct.call_args.args[0]['No'], 2)
        self.assertEqual(loaded.iloc[0]['Moon Semidiameter (deg)'], 0.27)

    def test_missing_legacy_timestamp_does_not_abort_remaining_rows(self):
        valid = self.raw_row(2)
        invalid = {**self.raw_row(3), 'Best_Time_Tel': None}
        for row in (valid, invalid):
            row.pop('moon_semidiameter_deg_BT')

        def reconstruct(row):
            if row['No'] == 3:
                raise ValueError('missing timestamp')
            return 0.27

        with patch.object(diagnostic, '_observation_semidiameter_deg', side_effect=reconstruct):
            loaded, output = self.load_rows([valid, invalid])
        self.assertEqual(loaded['No'].tolist(), [2])
        self.assertEqual(loaded.attrs['excluded_rows'][0]['No'], 3)
        self.assertIn('geometry_reconstruction', output)

    def test_all_invalid_inputs_return_an_empty_audited_frame(self):
        loaded, _ = self.load_rows([{**self.raw_row(), 'Status': 'invalid'}])
        self.assertTrue(loaded.empty)
        self.assertEqual(len(loaded.attrs['excluded_rows']), 1)

    def test_full_chain_rejects_invalid_weather_and_below_horizon_objects(self):
        loaded, _ = self.load_rows([self.raw_row()])
        row = loaded.iloc[0]
        for rh in (-1, 100, math.nan, math.inf):
            with self.subTest(rh=rh):
                with patch.object(diagnostic, 'hitung_sky_brightness') as sky:
                    with self.assertRaises(ValueError):
                        diagnostic.compute_full_chain(row, rh)
                    sky.assert_not_called()
        below_horizon = {**row.to_dict(), 'Moon Alt (°)': -1}
        with patch.object(diagnostic, 'hitung_sky_brightness') as sky:
            with self.assertRaises(ValueError):
                diagnostic.compute_full_chain(below_horizon, 75)
            sky.assert_not_called()
        # A real RH=1% must remain 1%, rather than become an assumed 5%.
        with patch.object(diagnostic, 'hitung_sky_brightness', return_value=dict(
            sky_brightness=1000, k_v=.2, extinction_mag_v=.5, transmission_v=10**-.2,
        )) as sky:
            result = diagnostic.compute_full_chain(row, 1)
        self.assertEqual(result['rh_used'], 1)
        self.assertEqual(sky.call_args.kwargs['humidity'], 1)

    def test_sensitivity_and_workbook_retain_ccd_scope_and_exclusion_audit(self):
        loaded, _ = self.load_rows([self.raw_row(2), {**self.raw_row(1), 'Status': 'invalid'}])
        with contextlib.redirect_stdout(io.StringIO()) as output:
            sensitivity, critical = diagnostic.analyze_rh_sensitivity(loaded)
            decomposition = diagnostic.analyze_error_decomposition(loaded)
            errors = diagnostic.analyze_error_bars(sensitivity, critical)
        self.assertIn('lintas metode', output.getvalue())
        self.assertNotIn('structural', output.getvalue())
        self.assertNotIn('correctable', output.getvalue())
        for result in (sensitivity, critical, decomposition, errors):
            self.assertGreater(len(result), 0)
            self.assertTrue(result['Observation_Method'].eq('ccd').all())
            self.assertTrue(result['Comparison_Scope'].eq('cross_method_descriptive').all())
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'diagnostic.xlsx'
            with contextlib.redirect_stdout(io.StringIO()):
                diagnostic.save_results(
                    sensitivity, critical, decomposition, errors, None, str(path),
                    input_audit={'fixture': loaded.attrs},
                )
            workbook = load_workbook(path)
            try:
                self.assertIn('Cakupan Perbandingan', workbook.sheetnames)
                self.assertIn('Input Dikeluarkan', workbook.sheetnames)
                headers = [cell.value for cell in workbook['Data Detail'][1]]
                self.assertIn('Observation_Method', headers)
                self.assertEqual(workbook['Input Dikeluarkan'].cell(2, 1).value, 'fixture')
                scope = dict(workbook['Cakupan Perbandingan'].values)
                self.assertIn('Tidak dilakukan', scope['Validasi visual empiris'])
            finally:
                workbook.close()


if __name__ == '__main__':
    unittest.main()
