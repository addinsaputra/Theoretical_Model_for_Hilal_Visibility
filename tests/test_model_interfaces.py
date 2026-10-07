"""Regresi kontrak Schaefer → Kastner → Crumey (tanpa API cuaca)."""

import math
import contextlib
import csv
import io
from datetime import date
from pathlib import Path
import sys
import unittest
import tempfile
from unittest.mock import patch


from hilal_visibility.models.geometry import crescent_area
from hilal_visibility.models.schaefer import hitung_sky_brightness, transmission_from_extinction_mag
from hilal_visibility.models.kastner import (
    S10_TO_NL, hitung_luminansi_intrinsik, terapkan_transmisi_atmosfer,
)
from hilal_visibility.models.crumey import (
    crescent_area_deg2, crescent_area_arcmin2, crescent_area_sr,
    crumey_visibility, is_visible, hilal_naked_eye_visibility,
    visibility_margin_mag, arcmin2_to_sr, nL_to_cdm2, mag_to_lux,
)
from hilal_visibility.calculator import HilalVisibilityCalculator
from hilal_visibility.reports.multi_location import _simpan_excel_multi, _plot_multi_lokasi
from hilal_visibility.batch import run_single_observation, save_to_csv, save_to_excel
from hilal_visibility.studies.diagnostics import compute_delta_m, compute_full_chain, load_observation_data
from openpyxl import load_workbook


class AtmosphericInterfaceTests(unittest.TestCase):
    def test_zero_and_one_mag_transmission(self):
        self.assertEqual(transmission_from_extinction_mag(0), 1)
        self.assertAlmostEqual(transmission_from_extinction_mag(1), 0.3981071705534972)

    def test_nonnegative_extinction_transmission_bounds(self):
        for extinction_mag in (0, 0.1, 1, 5, 20, 100):
            with self.subTest(extinction_mag=extinction_mag):
                transmission = transmission_from_extinction_mag(extinction_mag)
                self.assertGreater(transmission, 0)
                self.assertLessEqual(transmission, 1)

    def test_schaefer_keeps_coefficient_and_los_loss_separate(self):
        sky = hitung_sky_brightness(7, 2022, -1, 8, 75, 25, -7, 89, 5)
        self.assertEqual(len(sky['K']), 5)
        self.assertEqual(len(sky['DM']), 5)
        self.assertEqual(sky['k_v'], sky['K'][2])
        self.assertEqual(sky['extinction_mag_v'], sky['DM'][2])
        self.assertGreater(sky['extinction_mag_v'], sky['k_v'])
        self.assertAlmostEqual(sky['transmission_v'], 10 ** (-0.4 * sky['DM'][2]))

    def test_atmosphere_is_applied_once(self):
        # Sudut fase sengaja bukan 180 - elongasi untuk menguji pemisahan input.
        alpha, elongation, r = 170, 8, 0.26
        mv = 0.026 * alpha + 4e-9 * alpha**4 - 12.73
        area = 0.5 * math.pi * r**2 * (1 - math.cos(math.radians(elongation)))
        intrinsic = hitung_luminansi_intrinsik(alpha, elongation, r)
        # Mean luminance times solid angle must recover the phase-law flux,
        # using the same V zero point as Crumey's limiting magnitudes.
        flux = nL_to_cdm2(terapkan_transmisi_atmosfer(intrinsic, 1)) * area * math.radians(1)**2
        self.assertAlmostEqual(flux / mag_to_lux(mv), 1, places=12)
        self.assertEqual(terapkan_transmisi_atmosfer(intrinsic, 1), S10_TO_NL * intrinsic)
        transmission = transmission_from_extinction_mag(1)
        self.assertAlmostEqual(
            terapkan_transmisi_atmosfer(intrinsic, transmission) / (S10_TO_NL * intrinsic),
            transmission,
        )


class CrescentGeometryTests(unittest.TestCase):
    def test_area_conversions_and_shared_geometry(self):
        for elongation in (0, 8, 90, 180):
            with self.subTest(elongation=elongation):
                deg2 = crescent_area(elongation, 0.26)
                self.assertEqual(crescent_area_deg2(elongation, 0.26), deg2)
                self.assertEqual(crescent_area_arcmin2(elongation, 0.26), deg2 * 3600)
                self.assertEqual(crescent_area_sr(elongation, 0.26), arcmin2_to_sr(deg2 * 3600))
        self.assertEqual(crescent_area(0, 0.26), 0)
        self.assertAlmostEqual(crescent_area(180, 0.26), math.pi * 0.26**2)

    def test_zero_area_source(self):
        self.assertEqual(hitung_luminansi_intrinsik(180, 0, 0.26), 0)


class ContrastInterfaceTests(unittest.TestCase):
    def test_increment_contrast_in_all_crumey_entry_points(self):
        for ratio in (1, 0.1, 0):
            with self.subTest(ratio=ratio):
                self.assertAlmostEqual(crumey_visibility(ratio * 10, 10, 1e-6)['C_obj'], ratio)
                self.assertAlmostEqual(is_visible(ratio * 10, 10, 1e-6)['C_object'], ratio)
                result = hilal_naked_eye_visibility(ratio * 1000, 1000, 8, 0.26)
                self.assertAlmostEqual(result['C_obj'], ratio)

    def test_visibility_and_margin_have_same_sign(self):
        baseline = is_visible(1, 1, 1e-6)
        threshold = baseline['C_threshold']
        for factor in (0.1, 1, 10):
            result = is_visible(threshold * factor, 1, 1e-6)
            self.assertAlmostEqual(result['margin_mag'], 2.5 * math.log10(factor))
            self.assertEqual(result['visible'], factor > 1)

    def test_boundary_and_no_source(self):
        self.assertEqual(visibility_margin_mag(0.1, 0.1), 0)
        self.assertEqual(visibility_margin_mag(0, 0.1), float('-inf'))
        with patch('hilal_visibility.models.crumey.contrast_threshold', return_value=1):
            result = hilal_naked_eye_visibility(1000, 1000, 8, 0.26)
        self.assertEqual(result['delta_m'], 0)
        self.assertFalse(result['visible'])
        result = hilal_naked_eye_visibility(0, 1000, 8, 0.26)
        self.assertEqual(result['C_obj'], 0)
        self.assertEqual(result['delta_m'], float('-inf'))
        self.assertFalse(result['visible'])


class CoreInterfaceTests(unittest.TestCase):
    def setUp(self):
        self.calc = HilalVisibilityCalculator('Uji', -7, 110, 89, 'Asia/Jakarta', 9, 1444,
                                              sumber_atmosfer='manual')
        self.calc.hasil['tanggal_pengamatan'] = date(2023, 3, 22)
        self.posisi = dict(sun_alt=-1, sun_az=270, moon_alt=5, moon_az=278,
                           phase_angle=170, elongation=8, moon_semidiameter=0.26)

    def test_core_passes_explicit_schaefer_values(self):
        sky = dict(sky_brightness=1000, k_v=0.2, extinction_mag_v=2, transmission_v=10**-0.8)
        with patch('hilal_visibility.calculator.hitung_sky_brightness', return_value=sky):
            B, k, extinction, transmission = self.calc.hitung_sky_brightness_schaefer(75, 25, self.posisi)
        self.assertEqual((B, k, extinction, transmission), (1000, 0.2, 2, 10**-0.8))
        source = self.calc.hitung_luminansi_hilal_kastner(self.posisi, transmission)
        intrinsic = hitung_luminansi_intrinsik(170, 8, 0.26)
        self.assertAlmostEqual(source / intrinsic, S10_TO_NL * transmission)

    def test_core_does_not_substitute_assumed_extinction_on_error(self):
        with patch('hilal_visibility.calculator.hitung_sky_brightness', side_effect=ValueError('invalid')):
            with self.assertRaises(ValueError):
                self.calc.hitung_sky_brightness_schaefer(100, 25, self.posisi)

    def test_telescope_and_naked_eye_use_same_contrast(self):
        for ratio in (1, 0.1, 0):
            with self.subTest(ratio=ratio):
                C_ne, dm_ne, _ = self.calc.hitung_visibilitas_naked_eye(ratio * 1000, 1000, self.posisi)
                L_tel, B_tel, C_tel, _, dm_tel = self.calc.hitung_visibilitas_teleskop(
                    ratio * 1000, 1000, self.posisi)
                self.assertAlmostEqual(C_ne, ratio)
                self.assertAlmostEqual(C_tel, ratio)
                self.assertAlmostEqual(L_tel / B_tel, ratio)
                if ratio == 0:
                    self.assertEqual(dm_ne, float('-inf'))
                    self.assertEqual(dm_tel, float('-inf'))

    def test_zero_area_cannot_be_visible(self):
        self.posisi['elongation'] = 0
        _, dm_ne, result = self.calc.hitung_visibilitas_naked_eye(0, 1000, self.posisi)
        *_, dm_tel = self.calc.hitung_visibilitas_teleskop(0, 1000, self.posisi)
        self.assertEqual(result['A_sr'], 0)
        self.assertFalse(result['visible'])
        self.assertEqual(dm_ne, float('-inf'))
        self.assertEqual(dm_tel, float('-inf'))


class DiagnosticInterfaceTests(unittest.TestCase):
    def test_diagnostic_contrast_matches_core(self):
        for ratio in (1, 0.1, 0):
            for mode in ('naked_eye', 'telescope'):
                with self.subTest(ratio=ratio, mode=mode):
                    result = compute_delta_m(ratio * 1000, 1000, 8, mode, moon_sd_deg=0.26)
                    self.assertAlmostEqual(result['C_obj'], ratio)
                    if ratio == 0:
                        self.assertEqual(result['delta_m'], float('-inf'))

    def test_diagnostic_chain_uses_schaefer_transmission(self):
        row = {'Moon Alt (°)': 5, 'Phase Angle (°)': 170, 'Elongasi (°)': 8,
               'T (°C)': 25, 'Lat': -7, 'Elv': 89, 'Tanggal': '2023-03-22',
               'Sun Alt BT (°)': -1, 'Moon Semidiameter (deg)': 0.2637}
        result = compute_full_chain(row, 75)
        intrinsic = hitung_luminansi_intrinsik(170, 8, 0.2637)
        self.assertEqual(result['L_star_s10'], intrinsic)
        self.assertAlmostEqual(result['transmission_v'], 10 ** (-0.4 * result['extinction_mag_v']))
        self.assertAlmostEqual(result['L_hilal_nL'] / intrinsic, S10_TO_NL * result['transmission_v'])
        self.assertAlmostEqual(result['C_obj'], result['L_hilal_nL'] / result['B_sky_nL'])


class IntegrationAndExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Ephemeris DE440 asli dan atmosfer manual: sunset + scan waktu optimal.
        cls.calc = HilalVisibilityCalculator(
            'Uji Semarang', -6.917, 110.348, 89, 'Asia/Jakarta', 9, 1444,
            sumber_atmosfer='manual', manual_rh=75, manual_t=25,
        )
        with contextlib.redirect_stdout(io.StringIO()):
            cls.hasil = cls.calc.jalankan_perhitungan_lengkap(
                use_telescope=True, aperture=100, magnification=50,
                mode='optimal', interval_menit=5, min_moon_alt=2,
            )
        cls.shared = dict(mode='optimal', bulan_hijri=9, tahun_hijri=1444,
                          delta_day=0, sumber_atmosfer='manual', tel_params={}, F_naked=2.5)
        cls.multi = [dict(success=True, lokasi={'nama': 'Uji Semarang', 'lat': -6.917,
                                               'lon': 110.348, 'elv': 89}, hasil=cls.hasil)]
        obs = dict(no=1, nama='Uji Semarang', tanggal='2023-03-22', lat=-6.917,
                   lon=110.348, elv=89, bulan_hijri=9, tahun_hijri=1444,
                   bias_t=0, bias_rh=0, observed=True)
        cls.obs = obs
        with patch('hilal_visibility.batch.HilalVisibilityCalculator') as calculator:
            calculator.return_value.jalankan_perhitungan_lengkap.return_value = cls.hasil
            with contextlib.redirect_stdout(io.StringIO()):
                cls.batch = run_single_observation(obs, verbose=False)
        if not cls.batch['success']:
            raise AssertionError(cls.batch)

    def test_real_pipeline_sunset_and_timesteps(self):
        self.assertGreater(len(self.hasil['all_timestep_results']), 1)
        for result in [self.hasil] + self.hasil['all_timestep_results']:
            if result.get('valid') is False:
                continue
            self.assertAlmostEqual(result['transmission_v'], 10**(-0.4 * result['extinction_mag_v']))
            self.assertAlmostEqual(result['rasio_kontras_ne'],
                                   result['luminansi_hilal_nl'] / result['sky_brightness_nl'])
            self.assertAlmostEqual(result['rasio_kontras_tel'], result['rasio_kontras_ne'])
        for key in ('optimal_result_ne', 'optimal_result_tel'):
            self.assertIn('extinction_mag_v', self.hasil[key])
            self.assertIn('transmission_v', self.hasil[key])

    def test_single_excel_keeps_diagnostic_columns(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'single.xlsx'
            with contextlib.redirect_stdout(io.StringIO()):
                self.calc.simpan_ke_excel(str(path))
            wb = load_workbook(path)
            try:
                ws = wb['Timestep Data']
                headers = [cell.value for cell in ws[1]]
                for name, key in (('k_v (mag/airmass)', 'k_v'),
                                  ('extinction_mag_v (mag)', 'extinction_mag_v'),
                                  ('transmission_v', 'transmission_v')):
                    value = ws.cell(2, headers.index(name) + 1).value
                    expected = self.hasil['all_timestep_results'][0][key]
                    if key == 'transmission_v':
                        self.assertTrue(math.isclose(value, expected, rel_tol=1e-14))
                    else:
                        self.assertAlmostEqual(value, expected, delta=5e-5)
            finally:
                wb.close()

    def test_batch_csv_keeps_sunset_and_optimal_atmosphere(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'batch.csv'
            with contextlib.redirect_stdout(io.StringIO()):
                save_to_csv([self.batch], str(path))
            with path.open(encoding='utf-8-sig', newline='') as stream:
                row = next(csv.DictReader(stream))
            for suffix, result in (('Sunset', self.hasil), ('BT', self.hasil['optimal_result_tel']),
                                   ('NE_Optimal', self.hasil['optimal_result_ne'])):
                self.assertAlmostEqual(float(row[f'kV_{suffix}']), result['k_v'], delta=5e-5)
                for key in ('extinction_mag_v', 'transmission_v'):
                    self.assertEqual(float(row[f'{key}_{suffix}']), result[key])
                self.assertEqual(float(row[f'moon_semidiameter_deg_{suffix}']), result['moon_semidiameter'])
                self.assertEqual(float(row[f'moon_distance_km_{suffix}']), result['moon_distance_km'])
            self.assertEqual(float(row['Phase_Angle_BT']), self.hasil['optimal_result_tel']['phase_angle'])
            self.assertEqual(float(row['P_Sunset']), self.hasil['pressure'])
            self.assertEqual(float(row['P_NE_Optimal']), self.hasil['optimal_result_ne']['pressure'])
            self.assertEqual(float(row['P_Tel_Optimal']), self.hasil['optimal_result_tel']['pressure'])
            for suffix, mode, scene in (
                ('NE_Sunset', 'ne', self.hasil), ('Tel_Sunset', 'tel', self.hasil),
                ('NE_Optimal', 'ne', self.hasil['optimal_result_ne']),
                ('Tel_Optimal', 'tel', self.hasil['optimal_result_tel']),
            ):
                self.assertEqual(row[f'Regime_{suffix}'], scene[f'crumey_{mode}_regime'])
                self.assertEqual(row[f'Achromatic_Extrapolation_{suffix}'],
                                 str(scene[f'crumey_{mode}_achromatic_extrapolation']))
            self.assertEqual(float(row['Threshold_Difference_Sunset_mag']), self.hasil['threshold_difference_mag'])
            self.assertEqual(row['F_Comparison'], self.hasil['field_factor_comparison'])
            diagnostic = load_observation_data(str(path)).iloc[0]
            self.assertEqual(diagnostic['Moon Semidiameter (deg)'],
                             self.hasil['optimal_result_tel']['moon_semidiameter'])
            self.assertEqual(diagnostic['Phase Angle (°)'], self.hasil['optimal_result_tel']['phase_angle'])
            # CSV lama tanpa r masih bisa dianalisis: posisi dihitung dari timestamp BT.
            import pandas as pd
            old_csv = pd.read_csv(path)
            old_csv.drop(columns=['moon_semidiameter_deg_BT']).to_csv(path, index=False)
            reconstructed = load_observation_data(str(path)).iloc[0]['Moon Semidiameter (deg)']
            # CSV lama menyimpan HH:MM:SS, sedangkan waktu model memiliki mikrodetik.
            exported_time = self.hasil['optimal_time_tel'].replace(microsecond=0)
            expected = self.calc.hitung_posisi_matahari_bulan(exported_time)['moon_semidiameter']
            self.assertAlmostEqual(reconstructed, expected, places=10)

    def test_batch_excel_diagnostics_match_csv_values(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'batch.xlsx'
            with contextlib.redirect_stdout(io.StringIO()):
                save_to_excel([self.batch], str(path))
            wb = load_workbook(path)
            try:
                ws = wb['Hasil Observasi']
                self.assertEqual(ws.max_column, 60)
                visual = wb['Diagnostik Visual']
                diagnostics = dict(zip([cell.value for cell in visual[1]], [cell.value for cell in visual[2]]))
                self.assertEqual(diagnostics['Regime_NE_Sunset'], self.hasil['crumey_ne_regime'])
                self.assertEqual(diagnostics['Regime_Tel_Optimal'], self.hasil['optimal_result_tel']['crumey_tel_regime'])
                self.assertEqual(diagnostics['Achromatic_Extrapolation_NE_Sunset'],
                                 self.hasil['crumey_ne_achromatic_extrapolation'])
                self.assertAlmostEqual(ws.cell(3, 49).value, self.hasil['extinction_mag_v'])
                self.assertAlmostEqual(ws.cell(3, 50).value, self.hasil['transmission_v'])
                self.assertAlmostEqual(ws.cell(3, 53).value,
                                       self.hasil['optimal_result_tel']['extinction_mag_v'])
                self.assertAlmostEqual(ws.cell(3, 54).value,
                                       self.hasil['optimal_result_tel']['transmission_v'])
                self.assertTrue(math.isclose(ws.cell(3, 55).value, self.hasil['moon_semidiameter'],
                                             rel_tol=1e-14))
                self.assertAlmostEqual(ws.cell(3, 57).value,
                                       self.hasil['optimal_result_tel']['moon_semidiameter'])
                for ci, label, expected in (
                    (58, 'P_Sunset (hPa)', self.hasil['pressure']),
                    (59, 'P_NE_Optimal (hPa)', self.hasil['optimal_result_ne']['pressure']),
                    (60, 'P_Tel_Optimal (hPa)', self.hasil['optimal_result_tel']['pressure']),
                ):
                    self.assertEqual(ws.cell(2, ci).value, label)
                    self.assertAlmostEqual(ws.cell(3, ci).value, expected)
                    self.assertEqual(ws.cell(3, ci).number_format, '0.00')
            finally:
                wb.close()

    def test_batch_pressure_exports_keep_distinct_values_at_each_time(self):
        # Different values expose accidentally reusing sunset pressure at both optima.
        calculation = {
            **self.hasil,
            'pressure': 1001.123456,
            'optimal_result_ne': {**self.hasil['optimal_result_ne'], 'pressure': 1002.234567},
            'optimal_result_tel': {**self.hasil['optimal_result_tel'], 'pressure': 1003.345678},
        }
        with patch('hilal_visibility.batch.HilalVisibilityCalculator') as calculator:
            calculator.return_value.jalankan_perhitungan_lengkap.return_value = calculation
            with patch('hilal_visibility.batch.CALC_MODE', 'optimal'), contextlib.redirect_stdout(io.StringIO()):
                result = run_single_observation(self.obs, verbose=False)
        self.assertTrue(result['success'])
        self.assertEqual((result['pressure'], result['opt_ne_pressure'], result['opt_tel_pressure']),
                         (1001.123456, 1002.234567, 1003.345678))
        with tempfile.TemporaryDirectory() as folder:
            csv_path, xlsx_path = Path(folder) / 'pressure.csv', Path(folder) / 'pressure.xlsx'
            with contextlib.redirect_stdout(io.StringIO()):
                save_to_csv([result], str(csv_path))
                save_to_excel([result], str(xlsx_path))
            with csv_path.open(encoding='utf-8-sig', newline='') as stream:
                row = next(csv.DictReader(stream))
            self.assertEqual([float(row[name]) for name in ('P_Sunset', 'P_NE_Optimal', 'P_Tel_Optimal')],
                             [1001.123456, 1002.234567, 1003.345678])
            wb = load_workbook(xlsx_path)
            try:
                for ci, expected in zip((58, 59, 60), (1001.123456, 1002.234567, 1003.345678)):
                    self.assertAlmostEqual(wb['Hasil Observasi'].cell(3, ci).value, expected, places=10)
            finally:
                wb.close()

    def test_sunset_mode_leaves_optimal_pressures_blank(self):
        with patch('hilal_visibility.batch.HilalVisibilityCalculator') as calculator:
            calculator.return_value.jalankan_perhitungan_lengkap.return_value = self.hasil
            with patch('hilal_visibility.batch.CALC_MODE', 'sunset'), contextlib.redirect_stdout(io.StringIO()):
                result = run_single_observation(self.obs, verbose=False)
        self.assertTrue(result['success'])
        self.assertEqual(result['pressure'], self.hasil['pressure'])
        self.assertIsNone(result['opt_ne_pressure'])
        self.assertIsNone(result['opt_tel_pressure'])
        with tempfile.TemporaryDirectory() as folder:
            csv_path, xlsx_path = Path(folder) / 'sunset.csv', Path(folder) / 'sunset.xlsx'
            with contextlib.redirect_stdout(io.StringIO()):
                save_to_csv([result], str(csv_path))
                save_to_excel([result], str(xlsx_path))
            with csv_path.open(encoding='utf-8-sig', newline='') as stream:
                row = next(csv.DictReader(stream))
            self.assertEqual(float(row['P_Sunset']), self.hasil['pressure'])
            self.assertEqual(row['P_NE_Optimal'], '')
            self.assertEqual(row['P_Tel_Optimal'], '')
            wb = load_workbook(xlsx_path)
            try:
                self.assertIsNone(wb['Hasil Observasi'].cell(3, 59).value)
                self.assertIsNone(wb['Hasil Observasi'].cell(3, 60).value)
            finally:
                wb.close()

    def test_combined_excel_includes_optimal_diagnostics(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'multi.xlsx'
            with contextlib.redirect_stdout(io.StringIO()):
                _simpan_excel_multi(self.multi, self.shared, str(path))
            wb = load_workbook(path)
            try:
                ws = wb['Ringkasan Multi-Lokasi']
                headers = [cell.value for cell in ws[1]]
                for label, expected in (
                    ('extinction_mag_v (mag)', self.hasil['extinction_mag_v']),
                    ('transmission_v', self.hasil['transmission_v']),
                    ('transmission_v NE Optimal', self.hasil['optimal_result_ne']['transmission_v']),
                    ('transmission_v Tel Optimal', self.hasil['optimal_result_tel']['transmission_v']),
                ):
                    value = ws.cell(2, headers.index(label) + 1).value
                    if label.startswith('transmission_v'):
                        self.assertTrue(math.isclose(value, expected, rel_tol=1e-14))
                    else:
                        self.assertAlmostEqual(value, expected, delta=5e-5)
            finally:
                wb.close()

    def test_zero_source_margin_survives_export_and_plot(self):
        import matplotlib
        matplotlib.use('Agg')
        from copy import deepcopy
        calc = deepcopy(self.calc)
        for result in calc.hasil['all_timestep_results']:
            if result.get('valid'):
                result['delta_m_ne'] = result['delta_m_tel'] = float('-inf')
        multi = deepcopy(self.multi)
        multi[0]['hasil']['optimal_delta_m_ne'] = float('-inf')
        multi[0]['hasil']['optimal_delta_m_tel'] = float('-inf')
        with tempfile.TemporaryDirectory() as folder:
            xlsx = Path(folder) / 'zero.xlsx'
            with contextlib.redirect_stdout(io.StringIO()):
                calc.simpan_ke_excel(str(xlsx))
                self.assertTrue(calc.plot_visibility_margin(str(Path(folder) / 'zero.png')))
                self.assertTrue(_plot_multi_lokasi(multi, self.shared, str(Path(folder) / 'multi.png')))
            wb = load_workbook(xlsx)
            try:
                ws = wb['Timestep Data']
                headers = [cell.value for cell in ws[1]]
                margin_col = headers.index('Margin Teleskop') + 1
                self.assertEqual(ws.cell(2, margin_col).value, '-inf')
            finally:
                wb.close()


if __name__ == '__main__':
    unittest.main()
