"""Regression tests for IFS requests, hourly interpolation, and failure exports."""

import contextlib
import csv
from datetime import datetime, timedelta, timezone
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch

import numpy as np
import pandas as pd
import requests
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'Core'))

import atmosfer_ecmwf_ifs as weather
from atmosphere_provenance import atmosphere_audit_record, save_atmosphere_provenance
import core_crescent_visibility as core
import core_multi_location as batch


UTC = timezone.utc
WIB = timezone(timedelta(hours=7))
LOCATION = weather.ObservingLocation('Uji', -6.97, 110.42, 1258.0, 'Asia/Jakarta')


def sample_table(start='2024-04-09 10:00', periods=4):
    return pd.DataFrame({
        'date': pd.date_range(start, periods=periods, freq='h', tz='UTC'),
        'temperature_2m': [20.0 + i for i in range(periods)],
        'relative_humidity_2m': [80.0, 100.0, 60.0, 40.0][:periods] if periods <= 4 else 80.0,
        'surface_pressure': [1000.0 + 4 * i for i in range(periods)],
        'dew_point_2m': [18.0 + i for i in range(periods)],
    })


def api_response(table):
    hourly = Mock()
    times = table['date']
    hourly.Time.return_value = int(times.iloc[0].timestamp())
    hourly.TimeEnd.return_value = int((times.iloc[-1] + timedelta(hours=1)).timestamp())
    hourly.Interval.return_value = 3600
    variables = []
    for name in table.columns.drop('date'):
        variable = Mock()
        variable_id, unit_id, _ = weather._VARIABLE_SPECS[name]
        variable.Variable.return_value = variable_id
        variable.Unit.return_value = unit_id
        variable.Altitude.return_value = 0 if name == 'surface_pressure' else 2
        variable.ValuesAsNumpy.return_value = table[name].to_numpy(dtype=np.float32)
        variables.append(variable)
    hourly.VariablesLength.return_value = len(variables)
    hourly.Variables.side_effect = lambda i: variables[i]
    response = Mock()
    response.Model.return_value = weather.Model.ecmwf_ifs
    response.Hourly.return_value = hourly
    response.Latitude.return_value = -6.99
    response.Longitude.return_value = 110.43
    response.Elevation.return_value = LOCATION.altitude
    response.Timezone.return_value = b'GMT'
    response.TimezoneAbbreviation.return_value = b'GMT'
    response.UtcOffsetSeconds.return_value = 0
    return response


class RequestTests(unittest.TestCase):
    def fetch(self):
        return weather.fetch_weather_with_info(
            LOCATION.latitude, LOCATION.longitude, '2024-04-09', '2024-04-09',
            list(weather.ATMOSPHERIC_VARIABLES), 'UTC', elevation=LOCATION.altitude,
        )

    def test_pin_elevation_units_and_provenance(self):
        response = api_response(sample_table())
        client = Mock()
        client.weather_api.return_value = [response]
        with patch.object(weather, '_get_client', return_value=client):
            table, info = self.fetch()
        params = client.weather_api.call_args.kwargs['params']
        self.assertEqual(params['models'], 'ecmwf_ifs')
        self.assertEqual(params['elevation'], 1258)
        self.assertEqual(params['cell_selection'], 'land')
        self.assertEqual(params['temperature_unit'], 'celsius')
        self.assertEqual(client.weather_api.call_args.kwargs['timeout'], 30)
        self.assertEqual(str(table['date'].dt.tz), 'UTC')
        self.assertEqual(info['timezone'], 'GMT')
        self.assertEqual(info['elevation'], info['requested_elevation'])
        self.assertNotIn('grid_elevation', info)
        self.assertEqual(info['units']['surface_pressure'], 'hPa')
        self.assertEqual(info['model'], 'ecmwf_ifs')
        self.assertEqual(info['response_model_name'], 'ecmwf_ifs')
        self.assertEqual(info['response_model_id'], int(weather.Model.ecmwf_ifs))
        self.assertEqual(info['product_name'], weather.IFS_PRODUCT_NAME)
        self.assertEqual(info['product_type'], 'ifs_hres_hourly_time_series')
        self.assertIsNone(info['ifs_cycle'])
        self.assertIn('1cd0eaa', info['product_identity_reference'])
        self.assertEqual(table.attrs['weather_metadata'], info)

    def test_dataframe_wrapper_uses_shared_request_and_keeps_positional_arguments(self):
        frame = sample_table()
        with patch.object(weather, 'fetch_weather_with_info', return_value=(frame, {})) as fetch:
            actual = weather.fetch_hourly_weather(-7, 110, '2024-04-09', '2024-04-09',
                                                  ['temperature_2m'], 'UTC', elevation=89)
        self.assertIs(actual, frame)
        self.assertEqual(fetch.call_args.args[-1], 'UTC')
        self.assertEqual(fetch.call_args.kwargs['elevation'], 89)

    def test_transport_and_api_errors_preserve_cause(self):
        for error in (requests.ConnectionError('offline'), requests.Timeout('timeout'),
                      weather.openmeteo_requests.OpenMeteoRequestsError('bad request')):
            with self.subTest(error=type(error).__name__):
                client = Mock()
                client.weather_api.side_effect = error
                with patch.object(weather, '_get_client', return_value=client):
                    with self.assertRaises(weather.ECMWF_IFSAPIError) as raised:
                        self.fetch()
                self.assertIs(raised.exception.__cause__, error)

    def test_malformed_responses_are_domain_errors(self):
        for case in ('empty', 'multiple', 'model', 'hourly', 'interval', 'count',
                     'length', 'unit', 'variable', 'altitude', 'metadata'):
            with self.subTest(case=case):
                response = api_response(sample_table())
                responses = [response]
                hourly = response.Hourly()
                variable = hourly.Variables(0)
                if case == 'empty':
                    responses = []
                elif case == 'multiple':
                    responses *= 2
                elif case == 'model':
                    response.Model.return_value = 0
                elif case == 'hourly':
                    response.Hourly.return_value = None
                elif case == 'interval':
                    hourly.Interval.return_value = 7200
                elif case == 'count':
                    hourly.VariablesLength.return_value = 0
                elif case == 'length':
                    variable.ValuesAsNumpy.return_value = np.array([20.])
                elif case == 'unit':
                    variable.Unit.return_value = 999
                elif case == 'variable':
                    variable.Variable.return_value = weather.Variable.relative_humidity
                elif case == 'altitude':
                    variable.Altitude.return_value = 1000
                else:
                    response.Elevation.return_value = float('nan')
                client = Mock()
                client.weather_api.return_value = responses
                with patch.object(weather, '_get_client', return_value=client):
                    with self.assertRaises(weather.ECMWF_IFSAPIError):
                        self.fetch()

    def test_bad_request_inputs_do_not_contact_api(self):
        for overrides in ({'latitude': 91}, {'longitude': float('nan')},
                          {'elevation': float('inf')}, {'hourly_variables': []},
                          {'cell_selection': 'unknown'}, {'end_date': '2024-04-08'}):
            with self.subTest(overrides=overrides):
                kwargs = dict(latitude=-7, longitude=110, start_date='2024-04-09', end_date='2024-04-09')
                kwargs.update(overrides)
                with patch.object(weather, '_get_client') as client:
                    with self.assertRaises(ValueError):
                        weather.fetch_weather_with_info(**kwargs)
                    client.assert_not_called()


class WindowTests(unittest.TestCase):
    def setUp(self):
        self.table = sample_table()
        self.window = weather.AtmosphericWindow(self.table, {'model': 'ecmwf_ifs'})

    def test_exact_hour_midpoint_and_subseconds_use_native_anchors(self):
        exact = datetime(2024, 4, 9, 11, tzinfo=UTC)
        self.assertEqual(self.window.at_time(exact), (100, 21, 1004))
        self.assertEqual(self.window.at_time(exact + timedelta(minutes=30)), (80, 21.5, 1006))
        target = exact + timedelta(seconds=15, microseconds=500000)
        self.assertAlmostEqual(self.window.at_time(target)[0], 100 - 40 * 15.5 / 3600)

    def test_no_freezing_after_sunset_plus_one_hour(self):
        sunset = datetime(2024, 4, 9, 10, 30, tzinfo=UTC)
        self.assertNotEqual(self.window.at_time(sunset + timedelta(hours=1)),
                            self.window.at_time(sunset + timedelta(hours=1, minutes=30)))

    def test_sdk_second_resolution_accepts_microsecond_targets(self):
        table = self.table.copy()
        table['date'] = table['date'].astype('datetime64[s, UTC]')
        window = weather.AtmosphericWindow(table)
        target = datetime(2024, 4, 9, 10, 30, 15, 500000, tzinfo=UTC)
        self.assertEqual(window.at_time(target), self.window.at_time(target))

    def test_no_extrapolation_or_naive_times(self):
        for target in (datetime(2024, 4, 9, 9, 59, tzinfo=UTC),
                       datetime(2024, 4, 9, 13, 1, tzinfo=UTC)):
            with self.assertRaises(weather.ECMWF_IFSAPIError):
                self.window.at_time(target)
        with self.assertRaises(ValueError):
            self.window.at_time(datetime(2024, 4, 9, 11))

    def test_invalid_samples_are_rejected_without_clamping_raw_data(self):
        for column, value in (
            ('relative_humidity_2m', -1), ('relative_humidity_2m', 101),
            ('temperature_2m', float('inf')), ('temperature_2m', -273.15),
            ('surface_pressure', float('nan')), ('surface_pressure', 0),
            ('dew_point_2m', float('nan')),
        ):
            with self.subTest(column=column, value=value):
                table = self.table.copy()
                table.loc[1, column] = value
                with self.assertRaises(weather.ECMWF_IFSAPIError):
                    weather.AtmosphericWindow(table)

    def test_missing_duplicate_or_unsorted_hours_are_rejected(self):
        for table in (self.table.drop(index=1), self.table.iloc[[0, 1, 1, 2]],
                      self.table.iloc[::-1], self.table.drop(columns='surface_pressure')):
            with self.assertRaises(weather.ECMWF_IFSAPIError):
                weather.AtmosphericWindow(table)

    def test_midnight_fetches_two_dates_in_one_request(self):
        table = sample_table('2024-04-09 23:00', 2)
        target = datetime(2024, 4, 9, 23, 30, tzinfo=UTC)
        with patch.object(weather, 'fetch_weather_with_info', return_value=(table, {})) as fetch:
            self.assertEqual(weather.get_rh_t_at_time(LOCATION, target), (90, 20.5, 1002))
        self.assertEqual(fetch.call_count, 1)
        self.assertEqual(fetch.call_args.kwargs['start_date'], '2024-04-09')
        self.assertEqual(fetch.call_args.kwargs['end_date'], '2024-04-10')
        self.assertEqual(fetch.call_args.kwargs['elevation'], 1258)
        self.assertIn('dew_point_2m', fetch.call_args.kwargs['hourly_variables'])

    def test_wib_midnight_uses_previous_utc_date(self):
        table = sample_table('2024-04-09 17:00', 2)
        local = datetime(2024, 4, 10, 0, 30, tzinfo=WIB)
        with patch.object(weather, 'fetch_weather_with_info', return_value=(table, {})) as fetch:
            local_value = weather.get_rh_t_at_time(LOCATION, local)
            utc_value = weather.get_rh_t_at_time(LOCATION, local.astimezone(UTC))
        self.assertEqual(local_value, utc_value)
        self.assertEqual(fetch.call_args.kwargs['start_date'], '2024-04-09')
        self.assertEqual(fetch.call_args.kwargs['end_date'], '2024-04-09')

    def test_exact_last_hour_needs_no_next_day_or_unrelated_missing_samples(self):
        table = sample_table('2024-04-09 00:00', 24)
        table.loc[:22, 'temperature_2m'] = float('nan')
        target = datetime(2024, 4, 9, 23, tzinfo=UTC)
        with patch.object(weather, 'fetch_weather_with_info', return_value=(table, {})) as fetch:
            result = weather.get_rh_t_at_time(LOCATION, target)
        self.assertEqual(result, (80, 43, 1092))
        self.assertEqual(fetch.call_args.kwargs['end_date'], '2024-04-09')

    def test_missing_bounding_hour_is_not_extrapolated(self):
        target = datetime(2024, 4, 9, 10, 30, tzinfo=UTC)
        with patch.object(weather, 'fetch_weather_with_info', return_value=(self.table.iloc[:1], {})):
            with self.assertRaises(weather.ECMWF_IFSAPIError):
                weather.get_rh_t_at_time(LOCATION, target)

    def test_bias_validation_and_explicit_corrected_rh_clipping(self):
        self.assertEqual(weather.apply_bias_correction(95, 25, 1000, bias_t=-1, bias_rh=-10),
                         (100, 26, 1000))
        with self.assertRaises(weather.ECMWF_IFSAPIError):
            weather.apply_bias_correction(105, 25, 1000)
        with self.assertRaises(ValueError):
            weather.apply_bias_correction(95, 25, 1000, bias_t=float('nan'))


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.calc = core.HilalVisibilityCalculator(
            'Uji', -6.97, 110.42, 1258, 'Asia/Jakarta', 9, 1444, sumber_atmosfer='ecmwf_ifs',
        )

    def test_unsupported_source_fails_before_calculation(self):
        with patch.object(core, 'set_location') as location:
            with self.assertRaisesRegex(ValueError, 'Sumber atmosfer tidak dikenal'):
                core.HilalVisibilityCalculator(
                    'Uji', -7, 110, 89, 'Asia/Jakarta', 9, 1444, sumber_atmosfer='unknown',
                )
            location.assert_not_called()

    def test_source_menu_selects_manual_as_third_option(self):
        cases = (
            ([''], ('ecmwf_ifs', 80, 25, 1013.25)),
            (['2'], ('merra2', 80, 25, 1013.25)),
            (['3', '75', '22.5', '950'], ('manual', 75, 22.5, 950)),
        )
        for inputs, expected in cases:
            with self.subTest(inputs=inputs):
                with patch('builtins.input', side_effect=inputs):
                    with contextlib.redirect_stdout(io.StringIO()):
                        self.assertEqual(core._input_sumber_atmosfer(), expected)

    def test_api_failure_propagates_for_each_source(self):
        cases = (
            ('ecmwf_ifs', 'fetch_atmospheric_window', weather.ECMWF_IFSAPIError('offline')),
            ('merra2', 'merra2_get_rh_t', core.PowerAPIError('offline')),
        )
        for source, function, error in cases:
            with self.subTest(source=source):
                self.calc.sumber_atmosfer = source
                with patch.object(core, function, side_effect=error):
                    with self.assertRaises(type(error)):
                        self.calc._fetch_atmosfer(LOCATION, datetime(2024, 4, 9, 10, tzinfo=UTC),
                                                  verbose=False)
        self.assertNotIn('rh', self.calc.hasil)

    def test_manual_input_is_validated_and_bias_is_not_applied(self):
        self.calc.sumber_atmosfer = 'manual'
        self.calc.bias_t = 99
        self.calc.manual_rh, self.calc.manual_t, self.calc.manual_p = 75, 25, 880
        target = datetime(2024, 4, 9, 10, tzinfo=UTC)
        self.assertEqual(self.calc._fetch_atmosfer(LOCATION, target, verbose=False),
                         (75, 25, 75, 25, 880))
        self.assertEqual(self.calc.hasil['atmosphere_provenance'][0]['units']['temperature_2m'], '°C')
        self.calc.manual_rh = float('nan')
        with self.assertRaises(weather.ECMWF_IFSAPIError):
            self.calc._fetch_atmosfer(LOCATION, target, verbose=False)

    def test_raw_provenance_survives_correction(self):
        window = weather.AtmosphericWindow(sample_table(), {'model': 'ecmwf_ifs'})
        self.calc.bias_t, self.calc.bias_rh = -1, 5
        target = datetime(2024, 4, 9, 10, 30, tzinfo=UTC)
        with patch.object(core, 'fetch_atmospheric_window', return_value=window):
            result = self.calc._fetch_atmosfer(LOCATION, target, verbose=False)
        self.assertEqual(result, (90, 20.5, 85, 21.5, 1002))
        raw = self.calc.hasil['atmosphere_provenance'][0]['hourly_raw']
        self.assertEqual(raw[0]['relative_humidity_2m'], 80)
        self.assertIn('dew_point_2m', raw[0])

    def test_full_scan_and_refinement_use_hourly_samples_and_dynamic_coverage(self):
        sunset = datetime(2024, 4, 9, 17, 30, tzinfo=WIB)
        peak = sunset + timedelta(minutes=75)
        evaluated = []

        def fake_result(time_local, *args, cached_atm, **kwargs):
            evaluated.append((time_local, cached_atm))
            margin = 1 - abs((time_local - peak).total_seconds()) / 3600
            return dict(waktu_local=time_local, valid=True, moon_alt=5, sun_alt=-5,
                        delta_m_ne=margin, delta_m_tel=margin, rh=cached_atm[0],
                        temperature=cached_atm[1], sky_brightness_nl=1000,
                        luminansi_hilal_nl=100, telescope_gain=1)

        for interval in (1, 5):
            with self.subTest(interval=interval):
                evaluated.clear()
                table = sample_table('2024-04-09 10:00', 12)
                table.loc[:3, 'relative_humidity_2m'] = [80, 100, 60, 40]
                window = weather.AtmosphericWindow(table)
                with patch.object(core, 'fetch_atmospheric_window', return_value=window) as fetch:
                    with patch.object(self.calc, 'hitung_visibilitas_pada_waktu', side_effect=fake_result):
                        with contextlib.redirect_stdout(io.StringIO()):
                            result = self.calc.cari_visibilitas_optimal(
                                sunset, LOCATION, 100, 50, interval_menit=interval,
                            )
                self.assertEqual(fetch.call_count, 1)
                _, lower, upper = fetch.call_args.args
                self.assertEqual(lower, sunset.astimezone(UTC) - timedelta(minutes=1))
                self.assertEqual(upper, sunset.astimezone(UTC) + timedelta(minutes=3 + 119 * interval))
                self.assertEqual(result['total_timesteps'], 120)
                self.assertGreater(len(evaluated), 120)  # Includes refinement calls.
                for timestamp, values in evaluated:
                    self.assertEqual(values, window.at_time(timestamp))
                self.assertEqual(result['best_result_tel']['rh'], 70)


class FailureAndAuditExportTests(unittest.TestCase):
    def test_failed_weather_is_invalid_in_batch_and_exports(self):
        obs = dict(no=1, nama='Uji', tanggal='2024-04-09', lat=-6.97, lon=110.42, elv=1258,
                   bulan_hijri=10, tahun_hijri=1445, observed=False, bias_t=0, bias_rh=0)
        with patch.object(core, 'fetch_atmospheric_window', side_effect=weather.ECMWF_IFSAPIError('offline')):
            with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
                failed = batch.run_single_observation(obs, verbose=False)
        self.assertFalse(failed['success'])
        self.assertIn('offline', failed['error'])
        with tempfile.TemporaryDirectory() as folder:
            csv_path = Path(folder) / 'failed.csv'
            xlsx_path = Path(folder) / 'failed.xlsx'
            with contextlib.redirect_stdout(io.StringIO()):
                batch.save_to_csv([failed], str(csv_path))
                batch.save_to_excel([failed], str(xlsx_path))
            with csv_path.open(encoding='utf-8-sig', newline='') as stream:
                row = next(csv.DictReader(stream))
            self.assertEqual(row['Status'], 'invalid')
            self.assertEqual(row['Prediksi'], '')
            self.assertEqual(row['RH_Sunset'], '')
            self.assertEqual(row['Dm_Tel_BT'], '')
            for name in ('P_Sunset', 'P_NE_Optimal', 'P_Tel_Optimal'):
                self.assertEqual(row[name], '')
            wb = load_workbook(xlsx_path)
            try:
                sheet = wb['Hasil Observasi']
                self.assertEqual(sheet.cell(3, 48).value, 'ERROR')
                self.assertTrue(all(sheet.cell(3, i).value is None for i in range(9, 47)))
                self.assertTrue(all(sheet.cell(3, i).value is None for i in (58, 59, 60)))
            finally:
                wb.close()
            for path in (csv_path, xlsx_path):
                audit = json.loads(Path(str(path) + '.atmosphere.json').read_text(encoding='utf-8'))
                self.assertEqual(audit['observations'][0]['status'], 'invalid')
                self.assertIn('offline', audit['observations'][0]['error'])

    def test_raw_hourly_audit_serializes_losslessly_with_source_and_bias(self):
        window = weather.AtmosphericWindow(sample_table(), {
            'model': 'ecmwf_ifs', 'requested_elevation': 1258, 'elevation': 1258,
        })
        result = {'bias_t': -1, 'bias_rh': 5, 'atmosphere_provenance': [window.to_record()]}
        record = atmosphere_audit_record('Uji', 'ecmwf_ifs', True, result)
        with tempfile.TemporaryDirectory() as folder:
            path = save_atmosphere_provenance(Path(folder) / 'result.csv', [record])
            restored = json.loads(Path(path).read_text(encoding='utf-8'))
        audit = restored['observations'][0]
        self.assertEqual(audit['bias_t_celsius'], -1)
        self.assertEqual(audit['windows'][0], window.to_record())
        self.assertEqual(audit['windows'][0]['hourly_raw'][0]['date'], '2024-04-09T10:00:00+00:00')


if __name__ == '__main__':
    unittest.main()
