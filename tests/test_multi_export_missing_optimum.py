"""Regression for round(None) when the optimum is unavailable; no weather API."""

import contextlib
import csv
from datetime import datetime
import io
import math
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'Core'))
import core_multi_location as batch


def observation(no=1, **updates):
    time = datetime(2023, 3, 22, 18)
    row = dict(
        no=no, tanggal_obs='2023-03-22', nama='Export fixture', lat=-7, lon=110, elv=89,
        observed=False, success=True, bulan_hijri=9, tahun_hijri=1444,
        sunset_local=time, moon_alt_sunset=5, sun_alt_sunset=-1, elongation=8,
        moon_width=0.002, phase_angle=170, sky_brightness_nl=1000, luminansi_hilal_nl=10,
        k_v=0.2, rh=75, temperature=25, delta_m_ne_sunset=-1, delta_m_tel_sunset=-0.5,
        optimal_time_ne=time, optimal_time_tel=time, optimal_moon_alt_ne=4.123456,
        optimal_moon_alt_tel=4.234567, optimal_sun_alt_tel=-2,
        delta_m_ne_opt=-0.5, delta_m_tel_opt=-0.25,
        opt_ne_elongation=8, opt_ne_sky_brightness_nl=500, opt_ne_luminansi_hilal_nl=10,
        opt_ne_k_v=0.2, opt_ne_rh=75, opt_ne_temperature=25,
        opt_tel_elongation=8, opt_tel_sky_brightness_nl=400, opt_tel_luminansi_hilal_nl=10,
        opt_tel_k_v=0.2, opt_tel_rh=75, opt_tel_temperature=25, telescope_gain_opt=1.2,
        vis_duration_tel=0, pressure=1013.25, opt_ne_pressure=1012, opt_tel_pressure=1011,
    )
    row.update(updates)
    return row


class MissingOptimumExportTests(unittest.TestCase):
    def export(self, results):
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        xlsx, csv_path = Path(folder.name) / 'result.xlsx', Path(folder.name) / 'result.csv'
        with patch.object(batch, 'CALC_MODE', 'optimal'), contextlib.redirect_stdout(io.StringIO()):
            batch.save_to_excel(results, str(xlsx))
            batch.save_to_csv(results, str(csv_path))
            batch.print_results_table(results)
        wb = load_workbook(xlsx, data_only=True)
        self.addCleanup(wb.close)
        with csv_path.open(encoding='utf-8-sig', newline='') as stream:
            rows = list(csv.DictReader(stream))
        return wb, rows

    def test_missing_ne_altitude_and_optimum_do_not_abort_remaining_rows(self):
        wb, rows = self.export([
            observation(optimal_time_ne=None, optimal_moon_alt_ne=None, delta_m_ne_opt=None),
            observation(no=2, observed=True, delta_m_tel_opt=0.3),
        ])
        ws = wb['Hasil Observasi']
        self.assertTrue(all(ws.cell(3, col).value is None for col in range(24, 34)))
        self.assertEqual(ws.cell(3, 46).value, 'N')
        self.assertEqual(ws.cell(4, 46).value, 'Y')
        self.assertEqual(ws.cell(4, 48).value, '✓')
        self.assertAlmostEqual(ws.cell(4, 25).value, 4.1235)
        self.assertEqual(rows[0]['kV_NE_Optimal'], '')
        self.assertEqual(rows[1]['Prediksi'], 'Y')

    def test_missing_telescope_optimum_is_not_a_negative_prediction(self):
        # Real core can return -inf with no optimum; neither -inf nor its
        # placeholder zero geometry constitutes an evaluated optimal scene.
        wb, rows = self.export([
            observation(optimal_time_tel=None, optimal_moon_alt_tel=None,
                        optimal_sun_alt_tel=None, delta_m_tel_opt=-math.inf),
            observation(no=2),
        ])
        ws = wb['Hasil Observasi']
        self.assertTrue(all(ws.cell(3, col).value is None for col in range(34, 47)))
        self.assertIsNone(ws.cell(3, 48).value)
        self.assertEqual(rows[0]['Prediksi'], '')
        self.assertEqual(rows[0]['Dm_Tel_BT'], '')
        self.assertEqual(rows[0]['P_Tel_Optimal'], '')
        summary = {r[0].value: r[1].value for r in wb['Ringkasan']}
        self.assertEqual(summary['Observasi Berhasil'], 2)
        self.assertEqual(summary['Teleskop Dapat Dibandingkan'], 1)
        self.assertEqual(summary['Kecocokan Teleskop'], '1/1 (100.0%)')

    def test_none_optional_numeric_fields_and_failed_row_are_exportable(self):
        row = observation()
        for key in ('optimal_moon_alt_ne', 'opt_ne_elongation', 'opt_ne_sky_brightness_nl',
                    'opt_ne_luminansi_hilal_nl', 'opt_ne_k_v', 'opt_ne_rh', 'opt_ne_temperature',
                    'optimal_moon_alt_tel', 'optimal_sun_alt_tel', 'opt_tel_elongation',
                    'opt_tel_sky_brightness_nl', 'opt_tel_luminansi_hilal_nl', 'opt_tel_k_v',
                    'opt_tel_rh', 'opt_tel_temperature', 'telescope_gain_opt', 'vis_duration_tel'):
            row[key] = None
        # Exercise exporters directly; real atmospheric input validation remains
        # responsible for deciding whether a scene is physically valid.
        wb, rows = self.export([row, observation(no=2, success=False, error='weather unavailable')])
        ws = wb['Hasil Observasi']
        self.assertTrue(all(ws.cell(3, c).value is None for c in (*range(25, 32), *range(35, 45))))
        self.assertEqual(rows[0]['Moon_Alt_BT'], '')
        self.assertEqual(rows[0]['Leg_Time_min'], '')
        self.assertEqual(ws.cell(4, 48).value, 'ERROR')
        self.assertEqual(rows[1]['Status'], 'invalid')

    def test_genuine_minus_infinity_and_zero_are_preserved(self):
        wb, rows = self.export([observation(delta_m_tel_opt=-math.inf, vis_duration_tel=0)])
        ws = wb['Hasil Observasi']
        self.assertEqual(ws.cell(3, 45).value, '-inf')
        self.assertEqual(ws.cell(3, 46).value, 'N')
        self.assertEqual(ws.cell(3, 44).value, 0)
        self.assertEqual(rows[0]['Dm_Tel_BT'], '-inf')
        self.assertEqual(rows[0]['Leg_Time_min'], '0')


if __name__ == '__main__':
    unittest.main()
